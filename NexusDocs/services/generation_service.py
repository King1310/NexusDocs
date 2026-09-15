from __future__ import annotations

from datetime import timedelta
from pathlib import Path
import re
import os
import tempfile
import zipfile
from collections.abc import Iterable, Mapping

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
from lxml import etree

from domain.entities.organization import Organization
from domain.entities.person import Person
from domain.enums.organization_type import OrganizationType
from services.request_catalog import REQUEST_CATALOG, RequestTemplate
from utils.declension import decline_full_name, rank_genitive


class GenerationService:
    """Заполняет DOCX-шаблон данными человека и выбранной шапкой."""

    def generate(
        self,
        person: Person,
        organization: Organization,
        template_path: Path,
        output_dir: Path,
        overrides: Mapping[str, str] | None = None,
    ) -> Path:
        if not template_path.is_file():
            raise FileNotFoundError(f"Шаблон не найден: {template_path}")

        document = Document(template_path)
        replacements = self._build_replacements(person, organization, overrides)
        highlighted = self._highlighted_placeholders(organization)
        self._replace_in_container(document, replacements, highlighted)

        output_dir.mkdir(parents=True, exist_ok=True)
        filename = self._build_filename(person, organization)
        output_path = output_dir / filename
        document.save(output_path)
        self._replace_in_package(output_path, replacements, highlighted)
        return output_path

    def generate_bundle(
        self,
        person: Person,
        organizations: Iterable[Organization],
        template_dir: Path,
        output_dir: Path,
        overrides: Mapping[str, str] | None = None,
    ) -> list[Path]:
        """Create the complete 15-request bundle in catalogue order."""

        by_type = {
            organization.organization_type: organization
            for organization in organizations
        }
        generated: list[Path] = []
        for request in REQUEST_CATALOG:
            organization = (
                by_type.get(request.organization_type)
                if request.organization_type is not None
                else None
            )
            if request.requires_resolved_header and organization is None:
                raise ValueError(
                    f"Не найдена шапка: {request.organization_type.value}"
                )
            generated.append(
                self.generate_request(
                    person=person,
                    request=request,
                    organization=organization,
                    template_dir=template_dir,
                    output_dir=output_dir,
                    overrides=overrides,
                )
            )
        return generated

    def generate_request(
        self,
        person: Person,
        request: RequestTemplate,
        organization: Organization | None,
        template_dir: Path,
        output_dir: Path,
        overrides: Mapping[str, str] | None = None,
    ) -> Path:
        template_path = template_dir / request.filename
        if not template_path.is_file():
            raise FileNotFoundError(f"Шаблон не найден: {template_path}")

        document = Document(template_path)
        replacements = self._build_replacements(person, organization, overrides)
        if request.number in {10, 11}:
            replacements["предварительного следствия"] = (
                self._mvd_proceeding_stage(person)
            )
        highlighted = self._highlighted_placeholders(organization)
        self._replace_in_container(document, replacements, highlighted)

        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / self._build_request_filename(person, request)
        document.save(output_path)
        self._replace_in_package(output_path, replacements, highlighted)
        return output_path

    @staticmethod
    def _highlighted_placeholders(
        organization: Organization | None,
    ) -> frozenset[str]:
        """The final Word never uses automatic review highlighting."""

        return frozenset()

    @staticmethod
    def _build_replacements(
        person: Person,
        organization: Organization | None,
        overrides: Mapping[str, str] | None = None,
    ) -> dict[str, str]:
        passport = person.passport
        declined = decline_full_name(person.full_name)
        initials = (
            f"{person.full_name.first_name[:1]}."
            f"{person.full_name.middle_name[:1]}."
        )
        locality = (
            person.registration_address.locality
            or person.registration_address.city
            or person.registration_address.district
            or person.registration_address.region
        )
        passport_issued = ""
        if passport:
            passport_issued = (
                f"{passport.issued_by}, {passport.issue_date:%d.%m.%Y}"
            ).strip(", ")
        case_type = str(person.soch_case.case_type).strip()
        normalized_case_type = case_type.casefold().replace("ё", "е")
        is_criminal_case = (
            normalized_case_type == "уд"
            or "уголов" in normalized_case_type
        )
        case_number = re.sub(
            r"^\s*№\s*",
            "",
            person.soch_case.case_number or "",
        ).strip()
        case_location_phrase = (
            f"находится уголовное дело № {case_number}"
            if is_criminal_case
            else "находятся материалы проверки"
        )
        article_match = re.search(
            r"ч\.\s*(\d+(?:\.\d+)*)",
            person.soch_case.article,
        )
        article_part = article_match.group(1) if article_match else person.soch_case.article

        replacements = {
            # Legacy templates may contain either phrase as literal text.
            # These entries deliberately precede the placeholder replacement
            # so the freshly inserted phrase is not processed a second time.
            "находится уголовное дело": case_location_phrase,
            "находятся материалы проверки": case_location_phrase,
            "{{RECIPIENT_HEADER}}": (
                GenerationService._simple_recipient_header(organization)
                if organization
                else ""
            ),
            "{{FULL_NAME}}": person.full_name.full,
            "{{FULL_NAME_INITIALS}}": person.full_name.initials,
            "{{LAST_NAME}}": person.full_name.last_name,
            "{{FIRST_NAME}}": person.full_name.first_name,
            "{{MIDDLE_NAME}}": person.full_name.middle_name,
            "{{INITIALS}}": initials,
            "{{LAST_NAME_GENITIVE}}": declined.last_name_genitive,
            "{{FIRST_NAME_GENITIVE}}": declined.first_name_genitive,
            "{{MIDDLE_NAME_GENITIVE}}": declined.middle_name_genitive,
            "{{LAST_NAME_INSTRUMENTAL}}": declined.last_name_instrumental,
            "{{BIRTH_DATE}}": person.birth.birth_date.strftime("%d.%m.%Y"),
            "{{BIRTH_PLACE}}": person.birth.place.full,
            "{{REGISTRATION_ADDRESS}}": person.registration_address.full,
            "{{REGISTRATION_LOCALITY}}": locality,
            "{{PASSPORT_SERIES}}": passport.series if passport else "",
            "{{PASSPORT_NUMBER}}": passport.number if passport else "",
            "{{PASSPORT_ISSUED}}": passport_issued,
            "{{MILITARY_UNIT}}": GenerationService._military_unit_value(
                person.military.military_unit
            ),
            "{{MILITARY_RANK}}": person.military.rank,
            "{{MILITARY_RANK_GENITIVE}}": rank_genitive(person.military.rank),
            "{{MILITARY_POSITION}}": person.military.position,
            "{{MILITARY_ID}}": person.military.military_id,
            "{{SOCH_DATE}}": person.soch_case.soch_date.strftime("%d.%m.%Y"),
            "{{SOCH_PLACE}}": person.soch_case.soch_place,
            "{{CASE_TYPE}}": case_type,
            "{{CASE_LOCATION_PHRASE}}": case_location_phrase,
            "{{CASE_NUMBER}}": case_number,
            "{{ARTICLE}}": person.soch_case.article,
            "{{ARTICLE_PART}}": article_part,
            "{{OUTGOING_NUMBER}}": person.document_info.outgoing_number,
            "{{OUTGOING_DATE}}": person.document_info.outgoing_date.strftime(
                "%d.%m.%Y"
            ),
            "{{RESPONSE_DEADLINE}}": (
                person.document_info.outgoing_date + timedelta(days=10)
            ).strftime("%d.%m.%Y"),
        }
        for key, value in (overrides or {}).items():
            placeholder = key if key.startswith("{{") else "{{" + key + "}}"
            replacements[placeholder] = value
        return replacements

    @staticmethod
    def _military_unit_value(value: str) -> str:
        """Return a unit identifier without a duplicated textual prefix."""

        return re.sub(
            r"^\s*(?:(?:в\s*/\s*(?:ч(?:асть|асти)?|част(?:ь|и)))|"
            r"(?:войсков(?:ая|ой)\s+част(?:ь|и)))\s*(?:№\s*)?",
            "",
            value or "",
            flags=re.IGNORECASE,
        ).strip()

    @staticmethod
    def _mvd_proceeding_stage(person: Person) -> str:
        """Return the procedural stage used only in both МВД requests."""

        normalized = str(person.soch_case.case_type).casefold().replace("ё", "е")
        if normalized.strip() == "уд" or "уголов" in normalized:
            return "предварительного следствия"
        return "доследственной проверки"

    @classmethod
    def _simple_recipient_header(cls, organization: Organization) -> str:
        """Build a header using the fixed 14-header protocol."""

        recipient_lines = [
            cls._clean_header_line(line)
            for line in organization.recipient.splitlines()
        ]
        recipient_lines = [line for line in recipient_lines if line]

        cleaned_recipient: list[str] = []
        for index, line in enumerate(recipient_lines):
            if cls._is_personal_greeting(line):
                continue
            if cls._is_leader_name(line):
                continue
            if cls._is_rank_line(line):
                next_line = (
                    recipient_lines[index + 1]
                    if index + 1 < len(recipient_lines)
                    else ""
                )
                if cls._is_leader_name(next_line):
                    continue
            cleaned_recipient.append(line)

        cleaned_recipient = cls._canonical_recipient_lines(
            organization.organization_type,
            cleaned_recipient,
        )

        postal_address = cls._clean_postal_address(organization.postal_address)
        recipient_block = "\n".join(cleaned_recipient)
        contact_lines: list[str] = []
        if postal_address:
            contact_lines.append(postal_address)

        phones = tuple(
            line
            for phone in organization.phones[:2]
            if (line := cls._clean_header_line(phone))
        )
        if phones:
            contact_lines.append("тел.: " + ",\n".join(phones))
        if organization.email:
            email = cls._clean_header_line(organization.email)
            if email and "требует провер" not in email.casefold():
                contact_lines.append(f"эл. почта: {email}")

        if recipient_block and contact_lines:
            return recipient_block + "\n\n" + "\n".join(contact_lines)
        return recipient_block or "\n".join(contact_lines)

    @classmethod
    def _canonical_recipient_lines(
        cls,
        organization_type: OrganizationType,
        lines: list[str],
    ) -> list[str]:
        """Keep the nine variable headers in the same terse form as the samples."""

        if organization_type is OrganizationType.MILITARY_COMMISSARIAT:
            details = [
                cls._strip_repeated_military_organization(
                    line,
                    OrganizationType.MILITARY_COMMISSARIAT,
                )
                for line in lines
                if not re.match(
                    r"(?i)^военному комиссару$|^военного комиссариата$",
                    line,
                )
            ]
            details = [line for line in details if line]
            return ["Военному комиссару", "Военного комиссариата", *details]

        if organization_type is OrganizationType.MILITARY_COMMANDANT:
            details = [
                cls._sentence_case_if_upper(
                    cls._strip_repeated_military_organization(
                        line,
                        OrganizationType.MILITARY_COMMANDANT,
                    )
                )
                for line in lines
                if not re.match(
                    r"(?i)^военному коменданту$|^военной комендатуры$",
                    line,
                )
            ]
            details = [line for line in details if line]
            return ["Военному коменданту", "Военной комендатуры", *details]

        if organization_type is OrganizationType.TFOMS:
            region = cls._region_genitive(" ".join(lines))
            result = [
                "Директору",
                "Территориального фонда",
                "обязательного медицинского страхования",
            ]
            if region:
                result.append(region)
            return result

        if organization_type is OrganizationType.ELECTION_COMMISSION:
            detail = ""
            for line in lines:
                match = re.match(
                    r"(?i)^(?:тик|территориальн\w+\s+избирательн\w+\s+комисси\w+)\s*(.*)$",
                    line,
                )
                if match and match.group(1).strip():
                    detail = match.group(1).strip()
                    break
            if not detail:
                detail = next(
                    (
                        line
                        for line in lines
                        if not re.match(r"(?i)^председателю$", line)
                        and not re.match(
                            r"(?i)^территориальн\w+\s+избирательн\w+\s+комисси\w+$",
                            line,
                        )
                    ),
                    "",
                )
            result = [
                "Председателю",
                "Территориальной избирательной комиссии",
            ]
            if detail:
                result.append(detail)
            return result

        if organization_type in {
            OrganizationType.NARCOLOGY,
            OrganizationType.PSYCHIATRY,
            OrganizationType.HOSPITAL,
        }:
            details = [
                cls._strip_search_snippet(line)
                for line in lines
                if not re.match(r"(?i)^главному врачу$", line)
                and not re.match(r"(?i)^заведующему", line)
            ]
            required = {
                OrganizationType.NARCOLOGY: (
                    r"(?i)больниц|нарколог|психоневролог|гбуз|"
                    r"поликлиник|кабинет|диспансер"
                ),
                OrganizationType.PSYCHIATRY: (
                    r"(?i)больниц|психиатр|психоневролог|гбуз|"
                    r"поликлиник|кабинет|отделен|диспансер"
                ),
                OrganizationType.HOSPITAL: (
                    r"(?i)больниц|црб|цгб|гбуз|поликлиник"
                ),
            }[organization_type]
            details = [line for line in details if re.search(required, line)]
            details = [cls._abbreviate_medical_title(line) for line in details]
            return [
                "Главному врачу",
                *(details or ["медицинской организации по месту регистрации"]),
            ]

        if organization_type is OrganizationType.ADMINISTRATION:
            details = [
                line
                for line in lines
                if not re.match(r"(?i)^главе$", line)
                and not cls._is_full_name_line(line)
            ]
            details = [
                re.sub(r"(?i)^администраци[яи]\s+", "", line)
                for line in details
            ]
            return ["Главе", *details]

        if organization_type is OrganizationType.CONTRACT_SERVICE_POINT:
            region = cls._region_genitive(" ".join(lines))
            details = [
                line.strip(" \t\"«»")
                for line in lines
                if not re.match(r"(?i)^начальнику$", line)
                and not re.match(r"(?i)^пункта?\s+отбора$", line)
                and not re.match(r"(?i)^на\s+военную\s+службу\s+по\s+контракту$", line)
                and not re.match(
                    r"(?i)^пункт\s+отбора\s+на\s+военную\s+службу\s+по\s+контракту",
                    line,
                )
            ]
            return [
                "Начальнику",
                "Пункта отбора",
                "на военную службу по контракту",
                *( [region] if region else details ),
            ]

        return lines

    @staticmethod
    def _strip_repeated_military_organization(
        value: str,
        organization_type: OrganizationType,
    ) -> str:
        """Remove a repeated organization label but keep its territory."""

        patterns = {
            OrganizationType.MILITARY_COMMISSARIAT: (
                r"^(?:(?:объедин[её]нн(?:ый|ого)|межрайонн(?:ый|ого)|"
                r"районн(?:ый|ого)|городск(?:ой|ого))\s+)*"
                r"военн(?:ый|ого)\s+комиссариат(?:а)?\s*"
            ),
            OrganizationType.MILITARY_COMMANDANT: (
                r"^военн(?:ая|ой)\s+"
                r"(?:(?:автомобильн(?:ая|ой)|территориальн(?:ая|ой))\s+)*"
                r"комендатур(?:а|ы)\s*"
            ),
        }
        pattern = patterns.get(organization_type)
        if pattern is None:
            return value
        return re.sub(pattern, "", value, flags=re.IGNORECASE).strip(" ,")

    @staticmethod
    def _abbreviate_medical_title(value: str) -> str:
        """Use conventional medical abbreviations in dynamic Word headers."""

        replacements = (
            (r"\bгосударственное\s+бюджетное\s+учреждение\s+здравоохранения\b", "ГБУЗ"),
            (r"\bгосударственное\s+автономное\s+учреждение\s+здравоохранения\b", "ГАУЗ"),
            (r"\bгосударственное\s+каз[её]нное\s+учреждение\s+здравоохранения\b", "ГКУЗ"),
            (r"\bцентральная\s+районная\s+больница\b", "ЦРБ"),
            (r"\bцентральная\s+городская\s+больница\b", "ЦГБ"),
            (r"\bгородская\s+клиническая\s+больница\b", "ГКБ"),
            (r"\bклиническая\s+областная\s+больница\b", "КОБ"),
            (r"\bобластная\s+клиническая\s+больница\b", "ОКБ"),
        )
        result = value
        for pattern, abbreviation in replacements:
            result = re.sub(pattern, abbreviation, result, flags=re.IGNORECASE)
        return result

    @staticmethod
    def _sentence_case_if_upper(value: str) -> str:
        letters = "".join(character for character in value if character.isalpha())
        if letters and letters == letters.upper():
            return value[:1].upper() + value[1:].lower()
        return value

    @staticmethod
    def _strip_search_snippet(value: str) -> str:
        return re.split(
            r"(?i)\s+(?:напоминает|сообщает|информирует)\b|"
            r"\s+правила\s+гигиены\b|"
            r"\s+адрес\s*:|"
            r"\s+электронн\w*\s+почт\w*\s*:|"
            r"\s+головн\w*\s+организац\w*\s*:",
            value,
            maxsplit=1,
        )[0].strip(" \t,;:-")

    @staticmethod
    def _region_genitive(value: str) -> str:
        normalized = value.casefold()
        regions = {
            "московск": "Московской области",
            "ленинградск": "Ленинградской области",
            "краснодарск": "Краснодарского края",
            "ставропольск": "Ставропольского края",
            "иркутск": "Иркутской области",
        }
        for fragment, region in regions.items():
            if fragment in normalized:
                return region
        return ""

    @staticmethod
    def _clean_header_line(value: str) -> str:
        line = re.sub(r"\s+", " ", value).strip(" \t,;")
        line = re.split(r"\s*\|\s*", line, maxsplit=1)[0]
        line = re.sub(
            r"(?i)\s*[([]\s*(?:инн|огрн)\b[^)\]]*[)\]]",
            "",
            line,
        )
        line = re.split(
            r"(?i)\s+-\s+(?:сайт|официальный\s+сайт|телефон|адрес|"
            r"режим\s+работы)\b",
            line,
            maxsplit=1,
        )[0]
        line = re.split(r"(?:\.{3}|…)\s*", line, maxsplit=1)[0]
        line = re.sub(r"(?i)\s+(?:в|на|по)$", "", line)
        return line.strip()

    @classmethod
    def _clean_postal_address(cls, value: str) -> str:
        if "требует провер" in value.casefold():
            return ""
        address = cls._clean_header_line(value)
        address = re.split(
            r"(?i)\s*[.,;]?\s*(?:номер|дата\s+регистрации|пф|инн|огрн|"
            r"часы\s+работы|на\s+карте|реквизиты|координаты)\s*:",
            address,
            maxsplit=1,
        )[0]
        address = re.sub(
            r"^(?:\d{5,6}\s+)?\d{2,3}[.,]\d{4,}\s+",
            "",
            address,
        )
        address = re.sub(r"^\d{9,13}\)?\s*", "", address)
        address = re.split(
            r"(?i)\s+(?:посмотреть|часы\s+работы|показать\s+на\s+карте|"
            r"смотреть\s+на\s+карте)\b",
            address,
            maxsplit=1,
        )[0]
        if address.count("(") > address.count(")"):
            address = address.rsplit("(", maxsplit=1)[0]
        address = address.strip(" \t,;")
        has_index = bool(re.search(r"(?<!\d)\d{6}(?!\d)", address))
        has_street_number = bool(
            re.search(
                r"(?i)\b(?:ул\.?|улица|проспект|пр-т|шоссе|пер\.?|"
                r"переулок|площадь|наб\.?|набережная)\b[^\n,;]*\d+",
                address,
            )
        )
        if not has_index and not has_street_number:
            return ""
        return address

    @staticmethod
    def _is_leader_name(value: str) -> bool:
        return bool(
            re.fullmatch(
                r"(?:[А-ЯЁ]\.?\s*){1,2}[А-ЯЁ][а-яё-]+(?:ой|ому|у|а|ым|им)?"
                r"|[А-ЯЁ][а-яё-]+\s+(?:[А-ЯЁ]\.\s*){1,2}",
                value.strip(),
            )
        )

    @staticmethod
    def _is_rank_line(value: str) -> bool:
        return bool(
            re.fullmatch(
                r"(?i)(?:генерал|полковник|подполковник|майор|капитан|лейтенант)\S*"
                r"(?:\s+(?:полиции|юстиции|внутренней\s+службы))?",
                value.strip(),
            )
        )

    @staticmethod
    def _is_personal_greeting(value: str) -> bool:
        return bool(re.fullmatch(r"(?i)уважаем(?:ый|ая)\s+.+!", value.strip()))

    @staticmethod
    def _is_full_name_line(value: str) -> bool:
        return bool(
            re.fullmatch(
                r"[А-ЯЁ][а-яё-]+\s+[А-ЯЁ][а-яё-]+\s+[А-ЯЁ][а-яё-]+",
                value.strip(),
            )
        )

    def _replace_in_container(
        self,
        container: DocumentObject | _Cell,
        replacements: dict[str, str],
        highlighted_placeholders: frozenset[str] = frozenset(),
    ) -> None:
        for paragraph in container.paragraphs:
            self._replace_in_paragraph(
                paragraph,
                replacements,
                highlighted_placeholders,
            )

        for table in container.tables:
            self._replace_in_table(
                table,
                replacements,
                highlighted_placeholders,
            )

        if isinstance(container, DocumentObject):
            for section in container.sections:
                self._replace_in_container(
                    section.header,
                    replacements,
                    highlighted_placeholders,
                )
                self._replace_in_container(
                    section.footer,
                    replacements,
                    highlighted_placeholders,
                )

    def _replace_in_table(
        self,
        table: Table,
        replacements: dict[str, str],
        highlighted_placeholders: frozenset[str] = frozenset(),
    ) -> None:
        for row in table.rows:
            for cell in row.cells:
                self._replace_in_container(
                    cell,
                    replacements,
                    highlighted_placeholders,
                )

    @staticmethod
    def _apply_text_replacements(
        text: str,
        replacements: Mapping[str, str],
    ) -> str:
        """Apply template replacements without multiplying a case number."""

        replaced = text
        for placeholder, value in replacements.items():
            if placeholder == "находится уголовное дело":
                # The package-level pass sees text already processed by
                # python-docx. Once the number follows the phrase, leave it
                # untouched so every pass remains idempotent.
                replaced = re.sub(
                    r"находится уголовное дело(?!\s*№)",
                    value,
                    replaced,
                )
            else:
                replaced = replaced.replace(placeholder, value)
        return replaced

    @staticmethod
    def _replace_in_paragraph(
        paragraph: Paragraph,
        replacements: dict[str, str],
        highlighted_placeholders: frozenset[str] = frozenset(),
    ) -> None:
        for run in paragraph.runs:
            for placeholder, value in replacements.items():
                if placeholder in run.text:
                    replaced_run = GenerationService._apply_text_replacements(
                        run.text,
                        {placeholder: value},
                    )
                    if replaced_run == run.text:
                        continue
                    run.text = replaced_run
                    if placeholder == "{{RECIPIENT_HEADER}}":
                        GenerationService._format_recipient_run(run)
                    if placeholder in highlighted_placeholders:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    else:
                        run.font.highlight_color = None

        combined = "".join(run.text for run in paragraph.runs)
        if not any(placeholder in combined for placeholder in replacements):
            return

        should_highlight = any(
            placeholder in combined
            for placeholder in highlighted_placeholders
        )
        replaced = GenerationService._apply_text_replacements(
            combined,
            replacements,
        )

        if paragraph.runs:
            paragraph.runs[0].text = replaced
            if "{{RECIPIENT_HEADER}}" in combined:
                GenerationService._format_recipient_run(paragraph.runs[0])
            if should_highlight:
                paragraph.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
            else:
                paragraph.runs[0].font.highlight_color = None
            for run in paragraph.runs[1:]:
                run.text = ""

    @staticmethod
    def _format_recipient_run(run) -> None:
        """Use TNR while preserving the size authored in each template."""

        run.font.name = "Times New Roman"
        run_properties = run._element.get_or_add_rPr()
        run_fonts = run_properties.find(qn("w:rFonts"))
        if run_fonts is None:
            run_fonts = OxmlElement("w:rFonts")
            run_properties.insert(0, run_fonts)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            run_fonts.set(qn(f"w:{attribute}"), "Times New Roman")

    @staticmethod
    def _replace_in_package(
        document_path: Path,
        replacements: Mapping[str, str],
        highlighted_placeholders: frozenset[str] = frozenset(),
    ) -> None:
        """Replace slots inside legacy Word text boxes and drawing stories."""

        namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        paragraph_tag = f"{{{namespace}}}p"
        text_tag = f"{{{namespace}}}t"

        descriptor, temporary_name = tempfile.mkstemp(
            suffix=".docx", dir=document_path.parent
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            with zipfile.ZipFile(document_path, "r") as source, zipfile.ZipFile(
                temporary_path, "w", zipfile.ZIP_DEFLATED
            ) as target:
                for item in source.infolist():
                    data = source.read(item.filename)
                    if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                        root = etree.fromstring(data)
                        changed = False
                        for paragraph in root.iter(paragraph_tag):
                            nodes = []
                            for node in paragraph.iter(text_tag):
                                owner = node.getparent()
                                while owner is not None and owner.tag != paragraph_tag:
                                    owner = owner.getparent()
                                if owner is paragraph:
                                    nodes.append(node)
                            if not nodes:
                                continue
                            combined = "".join(node.text or "" for node in nodes)
                            should_highlight = any(
                                placeholder in combined
                                for placeholder in highlighted_placeholders
                            )
                            replaced = GenerationService._apply_text_replacements(
                                combined,
                                replacements,
                            )
                            if replaced == combined:
                                continue
                            is_recipient_header = "{{RECIPIENT_HEADER}}" in combined
                            first_node = nodes[0]
                            first_node.text = replaced
                            run = first_node.getparent()
                            run_properties = run.find(qn("w:rPr"))
                            if run_properties is not None:
                                for existing in run_properties.findall(
                                    qn("w:highlight")
                                ):
                                    run_properties.remove(existing)
                            if is_recipient_header:
                                if run_properties is None:
                                    run_properties = OxmlElement("w:rPr")
                                    run.insert(0, run_properties)
                                run_fonts = run_properties.find(qn("w:rFonts"))
                                if run_fonts is None:
                                    run_fonts = OxmlElement("w:rFonts")
                                    run_properties.insert(0, run_fonts)
                                for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                                    run_fonts.set(
                                        qn(f"w:{attribute}"),
                                        "Times New Roman",
                                    )
                            if should_highlight:
                                if run_properties is None:
                                    run_properties = OxmlElement("w:rPr")
                                    run.insert(0, run_properties)
                                highlight = OxmlElement("w:highlight")
                                highlight.set(qn("w:val"), "yellow")
                                run_properties.append(highlight)
                            for node in nodes[1:]:
                                node.text = ""
                            if "\n" in replaced:
                                parts = replaced.split("\n")
                                first_node.text = parts[0]
                                run = first_node.getparent()
                                insert_at = run.index(first_node) + 1
                                for part in parts[1:]:
                                    run.insert(
                                        insert_at,
                                        etree.Element(f"{{{namespace}}}br"),
                                    )
                                    insert_at += 1
                                    line_text = etree.Element(text_tag)
                                    line_text.set(
                                        "{http://www.w3.org/XML/1998/namespace}space",
                                        "preserve",
                                    )
                                    line_text.text = part
                                    run.insert(insert_at, line_text)
                                    insert_at += 1
                            changed = True
                        if changed:
                            data = etree.tostring(
                                root,
                                xml_declaration=True,
                                encoding="UTF-8",
                                standalone=True,
                            )
                    target.writestr(item, data)
            os.replace(temporary_path, document_path)
        finally:
            temporary_path.unlink(missing_ok=True)

    @staticmethod
    def _build_filename(
        person: Person,
        organization: Organization,
    ) -> str:
        base = (
            f"{person.full_name.last_name}"
            f"{person.full_name.first_name[:1]}"
            f"{person.full_name.middle_name[:1]}_"
            f"{person.document_info.outgoing_date:%d.%m.%Y}_"
            f"{organization.organization_type.slug}"
        )
        safe = re.sub(r"[<>:\"/\\|?*]", "_", base)
        return f"{safe}.docx"

    @staticmethod
    def _build_request_filename(
        person: Person,
        request: RequestTemplate,
    ) -> str:
        base = (
            f"{request.number:02d}_{request.slug}_"
            f"{person.full_name.last_name}"
            f"{person.full_name.first_name[:1]}"
            f"{person.full_name.middle_name[:1]}_"
            f"{person.document_info.outgoing_date:%d.%m.%Y}"
        )
        return re.sub(r"[<>:\"/\\|?*]", "_", base) + ".docx"
