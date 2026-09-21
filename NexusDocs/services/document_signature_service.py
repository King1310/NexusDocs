from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
from typing import TYPE_CHECKING

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.text.paragraph import Paragraph

if TYPE_CHECKING:
    from domain.value_objects.investigator import Investigator


class DocumentSignatureService:
    """Create an editable copy with the supplied investigator's scanned signature."""

    SIGNATURE_NAME_PREFIX = "NexusDocs signature "
    DEFAULT_ASSET_DIRECTORY = (
        Path(__file__).resolve().parents[1] / "assets" / "signatures"
    )

    def __init__(
        self,
        signature_path: Path | None = None,
        *,
        signature_directory: Path | None = None,
    ) -> None:
        """Optionally override one asset or the directory used by the catalog.

        ``signature_path`` exists primarily for focused tests and deployments
        which inject a single image. Normal application code resolves the
        selected investigator's ``signature_filename`` inside the shared asset
        directory.
        """
        self.signature_path_override = signature_path
        self.signature_directory = (
            signature_directory or self.DEFAULT_ASSET_DIRECTORY
        )

    def create_unsigned_copy(self, source: Path, destination: Path) -> int:
        """Remove signatory images, never letterhead artwork or source assets."""
        if source.resolve() == destination.resolve():
            raise ValueError("Исходный комплект и копия должны быть разными файлами.")
        document = Document(source)
        paragraphs = list(document.element.body.iter(qn("w:p")))
        removed = 0
        for index, element in enumerate(paragraphs):
            text = Paragraph(element, document).text
            if not re.search(r"О\s*\.\s*Р\s*\.\s*Агиров", text, re.IGNORECASE):
                continue
            if "юстиции" not in text.casefold():
                continue
            # In template 09 the image is anchored three paragraphs above
            # the name, inside the printed signatory's position block.
            block = [element]
            for previous in reversed(paragraphs[max(0, index - 5):index]):
                value = Paragraph(previous, document).text.strip().casefold()
                if value and not any(part in value for part in (
                    "руководитель", "военного следственного", "ск россии",
                )):
                    break
                block.append(previous)
            for paragraph in block:
                for tag in ("w:drawing", "w:pict"):
                    for drawing in list(paragraph.iter(qn(tag))):
                        drawing.getparent().remove(drawing)
                        removed += 1
        # Also handle any NexusDocs investigator signature already stamped in
        # a source file without changing its text. This includes documents
        # generated before executor selection was introduced.
        for drawing in list(document.element.body.iter(qn("w:drawing"))):
            if self._contains_generated_signature(drawing):
                drawing.getparent().remove(drawing)
                removed += 1
        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        return removed

    def create_signed_copy(
        self,
        source: Path,
        destination: Path,
        investigator: "Investigator",
    ) -> int:
        if source.resolve() == destination.resolve():
            raise ValueError("Исходный документ и итоговый комплект должны быть разными файлами.")
        signature_path = self._signature_path(investigator)
        if not signature_path.is_file():
            raise FileNotFoundError(
                f"Не найден файл подписи {investigator.initials_surname}:\n"
                f"{signature_path}"
            )

        document = Document(source)
        signer = self._signer_pattern(investigator.initials_surname)
        signature_name = self._signature_name(investigator)
        inserted = 0
        paragraphs = list(document.element.body.iter(qn("w:p")))
        for index, element in enumerate(paragraphs):
            paragraph = Paragraph(element, document)
            match = signer.search(paragraph.text)
            # Name in contact details or an addressee is not a signing location.
            if match is None:
                continue
            is_mvd_form = (
                paragraph.text.strip() == match.group().strip()
                and any("(заместитель" in Paragraph(previous, document).text
                        for previous in paragraphs[max(0, index - 3):index])
            )
            if "юстиции" not in paragraph.text.casefold() and not is_mvd_form:
                continue
            if self._contains_generated_signature(element):
                continue
            right_tabs = [tab.position for tab in paragraph.paragraph_format.tab_stops
                          if tab.alignment == WD_TAB_ALIGNMENT.RIGHT]
            text_width = (right_tabs[-1] if right_tabs else
                          document.sections[0].page_width - document.sections[0].left_margin
                          - document.sections[0].right_margin)
            self._insert_signature(
                paragraph,
                match.start(),
                signature_path=signature_path,
                signature_name=signature_name,
                text_width=text_width,
                compact=is_mvd_form,
                crowded_position=investigator.key == "sarkisyan",
            )
            inserted += 1

        if not inserted:
            raise ValueError(
                "В комплекте не найдены строки подписи "
                f"{investigator.initials_surname}."
            )
        destination.parent.mkdir(parents=True, exist_ok=True)
        document.save(destination)
        return inserted

    def _insert_signature(
        self,
        paragraph: Paragraph,
        name_offset: int,
        *,
        signature_path: Path,
        signature_name: str,
        text_width: int,
        compact: bool = False,
        crowded_position: bool = False,
    ) -> None:
        """Anchor the image to the signer's text without changing line height."""
        position = 0
        for run in paragraph.runs:
            if position + len(run.text) > name_offset:
                offset = name_offset - position
                if offset:
                    after = deepcopy(run._r)
                    run.text, tail = run.text[:offset], run.text[offset:]
                    # The name runs contain text only; preserve their formatting.
                    from docx.text.run import Run
                    after_run = Run(after, paragraph)
                    after_run.text = tail
                    run._r.addnext(after)
                    name_run = after
                else:
                    name_run = run._r
                break
            position += len(run.text)
        else:
            raise ValueError("Не удалось определить место подписи в строке.")

        framed = paragraph._p.find(qn("w:pPr") + "/" + qn("w:framePr")) is not None
        compact = compact or framed
        width = Mm(16 if compact else 34)
        picture_run = paragraph.add_run()
        shape = picture_run.add_picture(str(signature_path), width=width)
        # The compact МВД forms have narrow signature fields. Ordinary request
        # signatures must remain large enough to be clearly visible in print.
        max_height = Mm(10 if compact else (18 if crowded_position else 20))
        if shape.height > max_height:
            shape.width = round(shape.width * max_height / shape.height)
            shape.height = max_height
        width = shape.width
        name_run.addprevious(picture_run._r)
        inline = shape._inline

        anchor = OxmlElement("wp:anchor")
        for key, value in {
            "distT": "0", "distB": "0", "distL": "0", "distR": "0",
            "simplePos": "0", "relativeHeight": "251658240", "behindDoc": "0",
            "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
        }.items():
            anchor.set(key, value)
        simple = OxmlElement("wp:simplePos")
        simple.set("x", "0")
        simple.set("y", "0")
        anchor.append(simple)
        for axis, relative, offset in (
            ("H", "character" if compact else "margin",
             int(Mm(36)) if compact else int(text_width - width - Mm(32))),
            (
                "V",
                "line",
                (
                    -int(shape.height - Mm(3.5))
                    if compact
                    else -int(Mm(11 if crowded_position else 13))
                ),
            ),
        ):
            placement = OxmlElement(f"wp:position{axis}")
            placement.set("relativeFrom", relative)
            displacement = OxmlElement("wp:posOffset")
            displacement.text = str(offset)
            placement.append(displacement)
            anchor.append(placement)
        anchor.append(deepcopy(inline.find(qn("wp:extent"))))
        effect = OxmlElement("wp:effectExtent")
        for side in ("l", "t", "r", "b"):
            effect.set(side, "0")
        anchor.append(effect)
        anchor.append(OxmlElement("wp:wrapNone"))
        for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
            child = inline.find(qn(tag))
            if child is not None:
                anchor.append(deepcopy(child))
        anchor.find(qn("wp:docPr")).set("descr", signature_name)
        inline.getparent().replace(inline, anchor)

    def _signature_path(self, investigator: "Investigator") -> Path:
        if self.signature_path_override is not None:
            return self.signature_path_override
        # Catalog filenames are application-owned, but basename keeps an
        # accidental path component from escaping the signature directory.
        filename = Path(investigator.signature_filename).name
        return self.signature_directory / filename

    @classmethod
    def _signature_name(cls, investigator: "Investigator") -> str:
        return f"{cls.SIGNATURE_NAME_PREFIX}{investigator.key}"

    @classmethod
    def _contains_generated_signature(cls, element) -> bool:
        return any(
            (prop.get("descr") or "").startswith(cls.SIGNATURE_NAME_PREFIX)
            for prop in element.iter(qn("wp:docPr"))
        )

    @staticmethod
    def _signer_pattern(initials_surname: str) -> re.Pattern[str]:
        """Match a catalog name even when Word varies spaces around initials."""
        pattern = re.escape(initials_surname.strip())
        pattern = pattern.replace(r"\ ", r"\s*")
        pattern = pattern.replace(r"\.", r"\s*\.\s*")
        return re.compile(pattern, re.IGNORECASE)
