from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import _Cell


def _paragraphs(container: DocumentObject | _Cell):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs(cell)


def _replace_paragraph_text(paragraph, replacements: tuple[tuple[str, str], ...]):
    combined = "".join(run.text for run in paragraph.runs)
    if not combined:
        return

    if combined.startswith("Около 00 часов 00 минут 10.08.2026"):
        replaced = "{{SOCH_CIRCUMSTANCE_SENTENCE}}"
    else:
        replaced = combined
        for old, new in replacements:
            replaced = replaced.replace(old, new)
        replaced = replaced.replace(
            " добровольно обратился \nв ВСО",
            " добровольно обратился в ВСО",
        )

    if replaced == combined:
        return
    if not paragraph.runs:
        paragraph.add_run(replaced)
        return
    paragraph.runs[0].text = replaced
    for run in paragraph.runs[1:]:
        run.text = ""


def build_template(source: Path, destination: Path) -> Path:
    document = Document(source)
    replacements = (
        ("Около 15 часов 00 минут ", ""),
        ("«19» августа 2026г.", "{{THREE_DAY_LONG_COMPACT}}"),
        ("«19» августа 2026 г.", "{{THREE_DAY_LONG}}"),
        ("«26» августа 2026 г.", "{{TEN_DAY_LONG}}"),
        ("Шмелёва Владимира Юрьевича", "{{FULL_NAME_GENITIVE}}"),
        ("Шмелёва В.Ю.", "{{FULL_NAME_INITIALS_GENITIVE}}"),
        ("Шмелёв В.Ю.", "{{FULL_NAME_INITIALS}}"),
        ("заместитель командира взвода", "{{MILITARY_POSITION}}"),
        ("г. Ясиноватая ДНР", "{{MILITARY_DEPLOYMENT}}"),
        ("ч. 2.1 ст. 337 УК РФ", "{{ARTICLE}}"),
        ("ч. 5 ст. 337 УК РФ", "{{ARTICLE}}"),
        ("по контракту", "{{SERVICE_BASIS}}"),
        ("старшины", "{{MILITARY_RANK_GENITIVE}}"),
        ("старшина", "{{MILITARY_RANK}}"),
        ("10.08.2026", "{{SOCH_DATE}}"),
        ("17.08.2026", "{{REGISTRATION_DATE}}"),
        ("19.08.2026", "{{THREE_DAY_DATE}}"),
        ("26.08.2026", "{{TEN_DAY_DATE}}"),
        ("95375", "{{MILITARY_UNIT}}"),
    )

    for paragraph in _paragraphs(document):
        _replace_paragraph_text(paragraph, replacements)
    for section in document.sections:
        for paragraph in _paragraphs(section.header):
            _replace_paragraph_text(paragraph, replacements)
        for paragraph in _paragraphs(section.footer):
            _replace_paragraph_text(paragraph, replacements)

    split_date = document.tables[2].rows[0].cells
    _replace_paragraph_text(
        split_date[2].paragraphs[0],
        (("19", "{{THREE_DAY_DAY}}"),),
    )
    _replace_paragraph_text(
        split_date[4].paragraphs[0],
        (("августа", "{{THREE_DAY_MONTH}}"),),
    )
    _replace_paragraph_text(
        split_date[5].paragraphs[0],
        (("20", "{{THREE_DAY_YEAR_PREFIX}}"),),
    )
    _replace_paragraph_text(
        split_date[6].paragraphs[0],
        (("26", "{{THREE_DAY_YEAR_SUFFIX}}"),),
    )

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    print(build_template(args.source, args.destination))


if __name__ == "__main__":
    main()
