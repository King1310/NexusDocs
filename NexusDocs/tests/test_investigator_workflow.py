import os
from dataclasses import replace
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
import pytest
from docx import Document
from docx.oxml.ns import qn
from PySide6.QtWidgets import QApplication

from domain.value_objects.investigator import INVESTIGATORS, TSOMARTOV, LITVINOV, SARKISYAN
from database.repositories.person_repository import PersonRepository
from mappers.person_mapper import PersonMapper
from services.person_service import PersonService
from services.generation_service import GenerationService
from services.extension_generation_service import ExtensionGenerationService
from services.request_catalog import REQUEST_CATALOG
from tests.factories.person_factory import PersonFactory
from tests.test_database import reset_database
from ui.form_data import PersonFormData
from ui.person_investigator_window import PersonInvestigatorWindow
from ui.person_window import PersonWindow
from ui.form_navigation import BACK_DIALOG_CODE

ROOT = Path(__file__).resolve().parents[1]


def document_text(path):
    return "\n".join("".join(p.itertext()) for p in Document(path).element.body.iter(qn("w:t")))


def test_executor_survives_create_reload_and_edit():
    reset_database()
    repo = PersonRepository()
    form = PersonFormData.from_person(PersonFactory.create())
    form.investigator = LITVINOV
    person = PersonService(repo).save(form)
    assert repo.get_by_id(person.id).investigator == LITVINOV
    edited = PersonFormData.from_person(repo.get_by_id(person.id))
    assert edited.investigator == LITVINOV
    edited.investigator = SARKISYAN
    PersonService(repo).update(person.id, edited)
    assert repo.get_by_id(person.id).investigator == SARKISYAN
    legacy = PersonMapper.to_record(person)
    legacy.pop("investigator_key")
    assert PersonMapper.from_record(legacy).investigator == TSOMARTOV


def test_selector_and_back_preserve_choice_and_person_fields():
    app = QApplication.instance() or QApplication([])
    form = PersonFormData()
    selector = PersonInvestigatorWindow(form)
    selector.investigator_combo.setCurrentIndex(1)
    selector.save_form()
    details = PersonWindow(form)
    details.last_name_input.setText("Иванов")
    details.go_back()
    assert details.result() == BACK_DIALOG_CODE
    reopened = PersonInvestigatorWindow(form)
    assert reopened.investigator_combo.currentData() == "litvinov"
    assert PersonWindow(form).last_name_input.text() == "Иванов"


@pytest.mark.parametrize("profile", INVESTIGATORS, ids=lambda p: p.key)
def test_every_template_uses_selected_executor(tmp_path, profile):
    person = replace(PersonFactory.create(), investigator=profile)
    service = GenerationService()
    for request in REQUEST_CATALOG:
        path = service.generate_request(person, request, None, ROOT / "templates/requests", tmp_path)
        text = document_text(path)
        for cell in Document(path).element.body.iter(qn("w:tc")):
            assert cell.find(qn("w:p")) is not None, "Word cells require a paragraph"
        assert "{{INVESTIGATOR" not in text
        assert profile.initials_surname in text or profile.table_name in text
        assert profile.position_short in text
        for other in INVESTIGATORS:
            if other != profile:
                assert other.initials_surname not in text
                assert other.surname_initials not in text
        if request.number == 9:
            doc = Document(path)
            table_text = " ".join(t.text or "" for table in doc.tables for t in table._tbl.iter(qn("w:t")))
            assert profile.position_table in table_text
            assert profile.table_name in table_text
        if request.number == 11 and profile != TSOMARTOV:
            assert "(в звании)" not in text
    extension = ExtensionGenerationService().generate(
        person, ROOT / "templates/extensions/extension_to_10_days.docx", tmp_path
    )
    text = document_text(extension)
    assert "{{INVESTIGATOR" not in text
    assert profile.initials_surname in text
    assert "Саркисяна А.К." in text  # Fixed report author, not executor.
    if profile != TSOMARTOV:
        assert "Цомартов" not in text
