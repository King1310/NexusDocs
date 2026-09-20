from datetime import date
import json
from pathlib import Path
import subprocess

from docx import Document
from pypdf import PdfReader, PdfWriter
import pytest

from services.signed_scan_service import DispatchPackageStore, WordBundleHeaderEmails
from services.word_dispatch_service import (
    ReviewedWordConverter,
    WordDispatchService,
    WordExportResult,
)
from tests.factories.person_factory import PersonFactory


def _metadata():
    records = []
    page = 1
    for number in range(1, 16):
        records.append({
            "number": number,
            "page": page,
            "header": "",
            "text": (
                f"Главному врачу\rИсправленная организация {number}\r"
                + (f"эл. почта: edited{number}@example.ru\r" if number != 4 else "")
                + "В производстве военного следственного отдела находится дело.\r"
                + "Ответ направить на эл. почта: sender@example.ru"
            ),
            "footer": f"Исх. № Исх-5883/{number}-26\rот 11.09.2026",
        })
        page += 3 if number == 9 else 2 if number == 11 else 1
    return {"requests": records}, page - 1


def test_bookmarks_after_editing_define_all_pages_and_only_recipient_email():
    payload, total_pages = _metadata()
    requests = ReviewedWordConverter.requests_from_metadata(payload, total_pages)

    assert len(requests) == 15
    assert sum(request.page_count for request in requests) == total_pages
    assert requests[8].page_count == 3
    assert requests[10].page_count == 2
    assert requests[8].outgoing_number == "5883.9"
    assert all(request.outgoing_date == date(2026, 9, 11) for request in requests)
    assert requests[2].recipient_email == "edited3@example.ru"
    assert requests[3].recipient_email == ""
    assert all(request.recipient_email != "sender@example.ru" for request in requests)


def test_deleted_bookmark_and_shared_page_fail_before_splitting():
    payload, total_pages = _metadata()
    payload["requests"].pop(5)
    with pytest.raises(ValueError, match="границы запросов"):
        ReviewedWordConverter.requests_from_metadata(payload, total_pages)
    payload, total_pages = _metadata()
    payload["requests"][4]["page"] = payload["requests"][3]["page"]
    with pytest.raises(ValueError, match="разрыв страницы"):
        ReviewedWordConverter.requests_from_metadata(payload, total_pages)


def test_prepare_converts_selected_saved_edits_and_never_uses_db_email(tmp_path, monkeypatch):
    selected = tmp_path / "мой проверенный Word.docx"
    doc = Document()
    doc.add_paragraph("Ручные правки пользователя: новая организация")
    doc.save(selected)
    original_bytes = selected.read_bytes()
    calls = []

    class Converter:
        def convert(self, snapshot, workspace):
            calls.append(snapshot)
            assert snapshot != selected
            assert snapshot.read_bytes() == original_bytes
            assert "Ручные правки" in Document(snapshot).paragraphs[0].text
            payload, total_pages = _metadata()
            writer = PdfWriter()
            for page in range(total_pages):
                writer.add_blank_page(width=595 + page, height=842)
            pdf = workspace / "converted.pdf"
            with pdf.open("wb") as stream:
                writer.write(stream)
            return WordExportResult(
                pdf, ReviewedWordConverter.requests_from_metadata(payload, total_pages)
            )

    old_email_calls = []
    monkeypatch.setattr(
        WordBundleHeaderEmails, "emails_by_request",
        lambda *_: old_email_calls.append(True) or {4: "obsolete@example.ru"},
    )
    person = PersonFactory.create()
    service = WordDispatchService(converter=Converter())
    result = service.prepare(
        person=person, docx_path=selected, output_directory=tmp_path / "package",
    )

    assert len(calls) == 1
    assert selected.read_bytes() == original_bytes
    assert old_email_calls == []
    assert len(result.attachments) == 15
    assert sum(item.page_count for item in result.attachments) == 18
    assert len(PdfReader(result.attachments[8].path).pages) == 3
    assert result.attachments[3].recipient_email == ""
    assert result.attachments[2].recipient_email == "edited3@example.ru"
    assert result.attachments[2].path.name == "исх. № 5883.3 от 11.09.2026 комиссариат.pdf"
    assert len(DispatchPackageStore().load(result.manifest_path).attachments) == 15


def test_word_failure_and_timeout_do_not_create_package(tmp_path, monkeypatch):
    selected = tmp_path / "reviewed.docx"
    Document().save(selected)
    converter = ReviewedWordConverter(timeout_seconds=5)
    monkeypatch.setattr(converter, "_powershell", lambda: "powershell")
    stopped = []
    monkeypatch.setattr(converter, "_stop_own_word", lambda *args: stopped.append(args))

    def timeout(*args):
        raise subprocess.TimeoutExpired("powershell", 5)

    monkeypatch.setattr(converter, "_run_powershell", timeout)
    with pytest.raises(RuntimeError, match="не завершил преобразование"):
        WordDispatchService(converter=converter).prepare(
            person=PersonFactory.create(), docx_path=selected,
            output_directory=tmp_path / "package",
        )
    assert len(stopped) == 1
    assert not (tmp_path / "package").exists()
    assert selected.exists()


def test_converter_reads_fresh_pagination_from_word_export(tmp_path, monkeypatch):
    converter = ReviewedWordConverter()
    monkeypatch.setattr(converter, "_powershell", lambda: "powershell")

    def export(executable, script, env, timeout):
        writer = PdfWriter()
        payload, total_pages = _metadata()
        for _ in range(total_pages):
            writer.add_blank_page(width=595, height=842)
        with Path(env["NEXUSDOCS_WORD_PDF"]).open("wb") as stream:
            writer.write(stream)
        Path(env["NEXUSDOCS_WORD_METADATA"]).write_text(json.dumps(payload), encoding="utf-8")
        return subprocess.CompletedProcess([], 0)

    monkeypatch.setattr(converter, "_run_powershell", export)
    result = converter.convert(tmp_path / "saved.docx", tmp_path)
    assert result.requests[8].page_count == 3
    assert len(result.requests) == 15


def test_password_protected_or_invalid_docx_is_rejected_before_word(tmp_path):
    selected = tmp_path / "encrypted.docx"
    selected.write_bytes(b"not an unencrypted OOXML archive")
    class Converter:
        def convert(self, *args):
            pytest.fail("Word must not open password-protected input")
    with pytest.raises(ValueError, match="защищён паролем"):
        WordDispatchService(converter=Converter()).prepare(
            person=PersonFactory.create(), docx_path=selected,
            output_directory=tmp_path / "package",
        )
