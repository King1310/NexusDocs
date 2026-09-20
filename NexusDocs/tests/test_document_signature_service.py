from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
import base64
import pytest

from services.document_signature_service import DocumentSignatureService
from services.word_bundle_service import WordBundleService
from tests.factories.person_factory import PersonFactory


def investigator(
    key: str = "litvinov",
    initials_surname: str = "В.Н. Литвинов",
    signature_filename: str = "litvinov_vn.png",
):
    return SimpleNamespace(
        key=key,
        initials_surname=initials_surname,
        signature_filename=signature_filename,
    )


def test_signature_copy_preserves_text_and_only_signs_investigator(tmp_path):
    source = tmp_path / "unsigned.docx"
    signed = tmp_path / "signed.docx"
    document = Document()
    paragraph = document.add_paragraph("подполковник юстиции                 ")
    paragraph.add_run("В . Н . ")
    paragraph.add_run("Литвинов")
    document.add_paragraph("подполковник юстиции                 О.Р. Агиров")
    document.add_paragraph("Контакты: В.Н. Литвинов")
    document.save(source)
    original = source.read_bytes()
    asset = tmp_path / "test-signature.png"
    asset.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
    ))
    profile = investigator()
    count = DocumentSignatureService(asset).create_signed_copy(
        source, signed, profile
    )

    assert count == 1
    assert source.read_bytes() == original
    result = Document(signed)
    assert [p.text for p in result.paragraphs] == [p.text for p in document.paragraphs]
    anchors = list(result.element.body.iter(qn("wp:anchor")))
    assert len(anchors) == 1
    assert anchors[0].find(qn("wp:wrapNone")) is not None
    # Tall signature scans must not overlap the printed position above.
    assert int(anchors[0].find(qn("wp:extent")).get("cy")) <= Mm(10)
    assert anchors[0].find(qn("wp:positionH")).get("relativeFrom") == "margin"
    assert anchors[0].find(qn("wp:docPr")).get("descr") == (
        DocumentSignatureService.SIGNATURE_NAME_PREFIX + profile.key
    )
    assert not list(Document(source).element.body.iter(qn("wp:anchor")))


def test_missing_signature_keeps_unsigned_document_intact(tmp_path):
    source = tmp_path / "unsigned.docx"
    Document().save(source)
    original = source.read_bytes()
    with pytest.raises(FileNotFoundError, match="файл подписи"):
        DocumentSignatureService(tmp_path / "missing.png").create_signed_copy(
            source,
            tmp_path / "signed.docx",
            investigator(
                key="sarkisyan",
                initials_surname="А.К. Саркисян",
                signature_filename="sarkisyan_ak.png",
            ),
        )
    assert source.read_bytes() == original


def test_selected_investigator_resolves_own_asset_from_directory(tmp_path):
    source = tmp_path / "unsigned.docx"
    signed = tmp_path / "signed.docx"
    document = Document()
    document.add_paragraph("старший лейтенант юстиции       А.К. Саркисян")
    document.save(source)
    expected_asset = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
    )
    (tmp_path / "sarkisyan_ak.png").write_bytes(expected_asset)
    profile = investigator(
        key="sarkisyan",
        initials_surname="А.К. Саркисян",
        signature_filename="sarkisyan_ak.png",
    )

    assert DocumentSignatureService(
        signature_directory=tmp_path
    ).create_signed_copy(source, signed, profile) == 1

    with ZipFile(signed) as package:
        embedded = package.read("word/media/image1.png")
    assert embedded == expected_asset


def test_unsigned_copy_removes_generated_signature_for_any_investigator(tmp_path):
    asset = tmp_path / "signature.png"
    asset.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
    ))
    source = tmp_path / "source.docx"
    signed = tmp_path / "signed.docx"
    unsigned = tmp_path / "unsigned.docx"
    document = Document()
    document.add_paragraph("подполковник юстиции       В.Н. Литвинов")
    document.save(source)
    service = DocumentSignatureService(asset)
    service.create_signed_copy(source, signed, investigator())

    assert service.create_unsigned_copy(signed, unsigned) == 1
    assert not list(Document(unsigned).element.body.iter(qn("w:drawing")))
    assert Document(unsigned).paragraphs[0].text == document.paragraphs[0].text


def test_two_bundle_names_are_distinct_and_explicit():
    person = PersonFactory.create()
    unsigned = WordBundleService.bundle_filename(person, signed=False)
    signed = WordBundleService.bundle_filename(person, signed=True)
    assert unsigned.endswith("_без_подписи.docx")
    assert signed.endswith("_с_подписью.docx")
    assert signed != unsigned


def test_mvd_form_without_rank_gets_signature_but_contact_name_does_not(tmp_path):
    asset = tmp_path / "signature.png"
    asset.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
    ))
    source = tmp_path / "form.docx"
    destination = tmp_path / "signed.docx"
    doc = Document()
    doc.add_paragraph("В.Н. Литвинов")
    doc.add_paragraph("Заместитель руководителя")
    doc.add_paragraph("(заместитель)")
    doc.add_paragraph("ВСО СК России")
    doc.add_paragraph("                     В.Н. Литвинов")
    doc.save(source)
    assert DocumentSignatureService(asset).create_signed_copy(source, destination, investigator()) == 1
    result = Document(destination)
    assert not list(result.paragraphs[0]._p.iter(qn("wp:anchor")))
    assert len(list(result.paragraphs[-1]._p.iter(qn("wp:anchor")))) == 1


def test_unsigned_removes_agirov_signature_but_preserves_letterhead(tmp_path):
    asset = tmp_path / "image.png"
    asset.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
    ))
    source = tmp_path / "assembled.docx"
    target = tmp_path / "unsigned.docx"
    document = Document()
    document.add_paragraph("СК России").add_run().add_picture(str(asset))
    document.add_paragraph("Основной текст запроса")
    document.add_paragraph("Руководитель")
    document.add_paragraph("военного следственного отдела").add_run().add_picture(str(asset))
    document.add_paragraph()
    document.add_paragraph()
    document.add_paragraph("подполковник юстиции       О.Р. Агиров")
    document.save(source)
    original = source.read_bytes()

    assert DocumentSignatureService().create_unsigned_copy(source, target) == 1
    assert source.read_bytes() == original
    result = Document(target)
    assert len(list(result.element.body.iter(qn("w:drawing")))) == 1
    assert [p.text for p in result.paragraphs] == [p.text for p in document.paragraphs]
