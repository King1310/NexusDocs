from __future__ import annotations

from zipfile import ZipFile

import pytest
from docx import Document

from domain.enums.registration_mode import RegistrationMode
from services.request_catalog import REQUEST_CATALOG
from services.unique_request_generation_service import (
    UniqueRequestGenerationService,
)
from tests.factories.person_factory import PersonFactory


def _request(number: int):
    return next(item for item in REQUEST_CATALOG if item.number == number)


def _template_dir():
    from pathlib import Path

    return Path(__file__).resolve().parents[1] / "templates" / "requests"


def _package_text(path) -> str:
    with ZipFile(path) as package:
        return "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        ).replace("\u00a0", " ")


def test_required_fields_are_specific_to_selected_template():
    service = UniqueRequestGenerationService()

    mvd_fields = {
        field.key for field in service.required_subject_fields(_request(10))
    }
    commissariat_fields = {
        field.key for field in service.required_subject_fields(_request(3))
    }

    assert mvd_fields == {
        "full_name",
        "birth_date",
        "birth_place",
        "registration_address",
        "case_type",
        "case_number",
    }
    assert "passport" in commissariat_fields
    assert "registration_address" in commissariat_fields
    assert "soch_date" in commissariat_fields


def test_dynamic_template_requires_a_recipient(tmp_path):
    with pytest.raises(ValueError, match="необходимо указать шапку"):
        UniqueRequestGenerationService().generate(
            person=PersonFactory.create(),
            request=_request(3),
            template_dir=_template_dir(),
            output_dir=tmp_path,
        )


def test_fixed_two_page_template_keeps_all_content_and_replaces_header(tmp_path):
    custom_header = (
        "Заведующему отделом ЗАГС\n"
        "Главного управления ЗАГС тестового региона\n\n"
        "123456, Тестовый регион, г. Тестовый, ул. Новая, д. 1\n"
        "тел.: +7 (000) 111-22-33\n"
        "эл. почта: zags@example.test"
    )
    output = UniqueRequestGenerationService().generate(
        person=PersonFactory.create(),
        request=_request(2),
        template_dir=_template_dir(),
        output_dir=tmp_path,
        recipient_header=custom_header,
    )

    text = _package_text(output)
    assert text.count("Главного управления ЗАГС тестового региона") == 2
    assert "Одинцовским управлением" not in text
    assert "zags@example.test" in text


def test_register_later_blanks_top_fields_and_removes_every_footer(tmp_path):
    person = PersonFactory.create()
    output = UniqueRequestGenerationService().generate(
        person=person,
        request=_request(10),
        template_dir=_template_dir(),
        output_dir=tmp_path,
        registration_mode=RegistrationMode.LATER,
    )

    generated = Document(output)
    for section in generated.sections:
        footer_text = "\n".join(
            paragraph.text
            for paragraph in section.footer.paragraphs
        )
        footer_text += "\n".join(
            paragraph.text
            for table in section.footer.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        assert not footer_text.strip()

    package_text = _package_text(output)
    assert person.document_info.outgoing_number not in package_text
    assert person.document_info.outgoing_date.strftime("%d.%m.%Y") not in package_text
    assert "{{OUTGOING_" not in package_text


def test_mvd_selection_combines_both_covers_and_form(tmp_path):
    output = UniqueRequestGenerationService().generate(
        person=PersonFactory.create(),
        request=_request(10),
        template_dir=_template_dir(),
        output_dir=tmp_path,
        recipient_header=(
            "Начальнику\nМУ МВД России «Тестовое»\n\n"
            "123456, г. Тестовый, ул. Проверочная, д. 1"
        ),
    )

    document = Document(output)
    text = _package_text(output)
    assert len(document.sections) >= 2
    assert "МУ МВД России «Тестовое»" in text
    # The custom local cover no longer addresses the previous chief by name;
    # the one remaining greeting belongs to the unchanged regional cover.
    assert text.count("Уважаемый Вячеслав Николаевич") == 1
    assert "{{" not in text


def test_unique_generation_does_not_modify_source_person(tmp_path):
    person = PersonFactory.create()
    original_document_info = person.document_info

    UniqueRequestGenerationService().generate(
        person=person,
        request=_request(1),
        template_dir=_template_dir(),
        output_dir=tmp_path,
    )

    assert person.document_info is original_document_info
