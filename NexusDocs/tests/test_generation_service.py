from dataclasses import replace
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt

from domain.entities.organization import Organization
from domain.enums.organization_type import OrganizationType
from domain.value_objects.address import Address
from services.generation_service import GenerationService
from services.request_catalog import REQUEST_CATALOG
from services.word_bundle_service import WordBundleService
from tests.factories.person_factory import PersonFactory


def test_military_unit_prefix_is_removed_without_leaving_word_fragments():
    for value in (
        "в/ч 95375",
        "в/часть 95375",
        "в/части 95375",
        "войсковая часть № 95375",
    ):
        assert GenerationService._military_unit_value(value) == "95375"


def test_generation_replaces_template_placeholders(tmp_path):
    template_path = tmp_path / "request_template.docx"
    template = Document()
    template.add_paragraph("{{RECIPIENT_HEADER}}")
    template.add_paragraph("Исх. № {{OUTGOING_NUMBER}} от {{OUTGOING_DATE}}")
    template.add_paragraph("В отношении: {{FULL_NAME}}, {{BIRTH_DATE}} г.р.")
    template.add_paragraph("Адрес: {{REGISTRATION_ADDRESS}}")
    template.save(template_path)

    organization = Organization(
        id=1,
        organization_type=OrganizationType.MILITARY_COMMISSARIAT,
        recipient="Военному комиссару\nВоенного комиссариата",
        postal_address="140100, Московская область, г. Раменское",
        phones=("+7 (496) 000-00-00",),
        email="example@example.ru",
        verification_status="test",
    )

    output_path = GenerationService().generate(
        person=PersonFactory.create(),
        organization=organization,
        template_path=template_path,
        output_dir=tmp_path / "output",
    )

    generated = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in generated.paragraphs)

    assert output_path.name.startswith("ИвановИИ_")
    assert "{{" not in text
    assert "Военному комиссару" in text
    assert "Иванов Иван Иванович" in text
    assert "123/1" in text


def test_ud_case_phrase_contains_one_number_sign_and_case_number():
    person = PersonFactory.create()
    person.soch_case = replace(
        person.soch_case,
        case_type="УД",
        case_number="№ 1.26.0200.2402.000152",
    )

    replacements = GenerationService._build_replacements(person, None)

    assert replacements["{{CASE_LOCATION_PHRASE}}"] == (
        "находится уголовное дело № 1.26.0200.2402.000152"
    )
    assert replacements["{{CASE_NUMBER}}"] == "1.26.0200.2402.000152"


def test_empty_house_is_omitted_from_address():
    address = Address(
        region="Вологодская область",
        district="",
        city="Белозерск",
        locality="",
        street="",
        house="",
    )

    assert address.full == "Вологодская область, Белозерск"
    assert "д." not in address.full


def test_military_unit_prefix_is_not_duplicated_by_templates():
    person = PersonFactory.create()
    person.military = replace(person.military, military_unit="в/ч 52033")

    replacements = GenerationService._build_replacements(person, None)

    assert replacements["{{MILITARY_UNIT}}"] == "52033"


def test_compound_article_part_is_kept_in_full():
    person = PersonFactory.create()
    person.soch_case = replace(
        person.soch_case,
        article="ч. 2.1 ст. 337 УК РФ",
    )

    replacements = GenerationService._build_replacements(person, None)

    assert replacements["{{ARTICLE_PART}}"] == "2.1"


def test_response_deadline_is_ten_days_after_outgoing_date():
    person = PersonFactory.create()
    person.document_info = replace(
        person.document_info,
        outgoing_date=date(2026, 4, 10),
    )

    replacements = GenerationService._build_replacements(person, None)

    assert replacements["{{RESPONSE_DEADLINE}}"] == "20.04.2026"


def test_response_deadline_keeps_bold_template_formatting(tmp_path):
    template_path = tmp_path / "deadline_template.docx"
    template = Document()
    paragraph = template.add_paragraph("Исполнить запрос в срок до ")
    deadline_run = paragraph.add_run("{{RESPONSE_DEADLINE}}")
    deadline_run.bold = True
    template.save(template_path)

    person = PersonFactory.create()
    person.document_info = replace(
        person.document_info,
        outgoing_date=date(2026, 4, 10),
    )
    organization = Organization(
        id=1,
        organization_type=OrganizationType.MILITARY_COMMISSARIAT,
        recipient="Военному комиссару",
        postal_address="140100, Московская область, г. Раменское",
    )

    output_path = GenerationService().generate(
        person=person,
        organization=organization,
        template_path=template_path,
        output_dir=tmp_path / "output",
    )
    generated = Document(output_path)
    deadline_runs = [
        run
        for paragraph in generated.paragraphs
        for run in paragraph.runs
        if "20.04.2026" in run.text
    ]

    assert len(deadline_runs) == 1
    assert deadline_runs[0].bold is True


def test_real_templates_use_bold_response_deadline_placeholder():
    template_dir = (
        Path(__file__).resolve().parents[1] / "templates" / "requests"
    )
    legacy_phrases = (
        "в кратчайший срок",
        "в максимально короткий срок",
        "исполнить запрос незамедлительно",
    )
    deadline_nodes = []

    for template_path in sorted(template_dir.glob("*.docx")):
        if template_path.name.startswith("~$"):
            continue
        document = Document(template_path)
        combined = "".join(
            node.text or ""
            for node in document.element.body.iter()
            if node.tag.endswith("}t")
        )
        assert not any(phrase in combined for phrase in legacy_phrases)
        for node in document.element.body.iter(qn("w:t")):
            if node.text == "{{RESPONSE_DEADLINE}}":
                deadline_nodes.append(node)
                run_properties = node.getparent().find(qn("w:rPr"))
                assert run_properties is not None
                assert run_properties.find(qn("w:b")) is not None

    assert len(deadline_nodes) == 13


def test_mp_case_phrase_remains_materials_check():
    person = PersonFactory.create()
    person.soch_case = replace(
        person.soch_case,
        case_type="МП",
        case_number=None,
    )

    replacements = GenerationService._build_replacements(person, None)

    assert replacements["{{CASE_LOCATION_PHRASE}}"] == (
        "находятся материалы проверки"
    )


def test_mvd_requests_change_proceeding_stage_for_mp_and_ud(tmp_path):
    template_dir = (
        Path(__file__).resolve().parents[1] / "templates" / "requests"
    )
    mvd_requests = REQUEST_CATALOG[9:11]

    for case_type, expected, absent in (
        ("МП", "доследственной проверки", "предварительного следствия"),
        ("УД", "предварительного следствия", "доследственной проверки"),
    ):
        person = PersonFactory.create()
        person.soch_case = replace(
            person.soch_case,
            case_type=case_type,
            case_number=("1.26.0200.2402.000152" if case_type == "УД" else None),
        )
        for request in mvd_requests:
            output = GenerationService().generate_request(
                person=person,
                request=request,
                organization=None,
                template_dir=template_dir,
                output_dir=tmp_path / case_type,
            )
            generated = Document(output)
            text = "".join(
                node.text or ""
                for node in generated.element.body.iter(qn("w:t"))
            )
            assert expected in text
            assert absent not in text


def test_ud_phrase_is_applied_to_every_real_request_template(tmp_path):
    person = PersonFactory.create()
    person.soch_case = replace(
        person.soch_case,
        case_type="УД",
        case_number="№ 1.26.0200.2402.000152",
    )
    organizations = [
        Organization(
            id=index,
            organization_type=organization_type,
            recipient=f"Адресат {index}",
            postal_address=f"14010{index}, Московская область, адрес {index}",
            phones=("+7 (495) 111-11-11",),
            email="example@example.ru",
            verification_status="test",
        )
        for index, organization_type in enumerate(OrganizationType, start=1)
    ]
    template_dir = (
        Path(__file__).resolve().parents[1] / "templates" / "requests"
    )

    generated = GenerationService().generate_bundle(
        person=person,
        organizations=organizations,
        template_dir=template_dir,
        output_dir=tmp_path / "output",
    )

    texts = []
    for path in generated:
        document = Document(path)
        texts.append("".join(
            node.text or ""
            for node in document.element.body.iter()
            if node.tag.endswith("}t")
        ))
    combined = "\n".join(texts)
    expected = "находится уголовное дело № 1.26.0200.2402.000152"
    assert "находятся материалы проверки" not in combined
    assert combined.count(expected) == 15
    assert combined.count("1.26.0200.2402.000152") == 15
    assert "1.26.0200.2402.000152 №" not in combined


def test_generation_never_highlights_automatic_header(tmp_path):
    template_path = tmp_path / "review_template.docx"
    template = Document()
    template.add_paragraph("{{RECIPIENT_HEADER}}")
    template.save(template_path)

    organization = Organization(
        id=None,
        organization_type=OrganizationType.HOSPITAL,
        recipient="Главному врачу\nГБУЗ «ЦРБ»",
        postal_address="ПОЧТОВЫЙ АДРЕС ТРЕБУЕТ ПРОВЕРКИ",
        email="ТРЕБУЕТСЯ ПРОВЕРКИ",
        verification_status="needs_review",
    )

    output_path = GenerationService().generate(
        person=PersonFactory.create(),
        organization=organization,
        template_path=template_path,
        output_dir=tmp_path / "output",
    )

    generated = Document(output_path)
    header_runs = [
        run
        for paragraph in generated.paragraphs
        for run in paragraph.runs
        if "Главному врачу" in run.text
    ]
    assert header_runs
    assert header_runs[0].font.highlight_color is None


def test_generation_clears_template_highlight_for_user_verified_header(tmp_path):
    template_path = tmp_path / "verified_template.docx"
    template = Document()
    paragraph = template.add_paragraph()
    run = paragraph.add_run("{{RECIPIENT_HEADER}}")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    template.save(template_path)

    organization = Organization(
        id=None,
        organization_type=OrganizationType.HOSPITAL,
        recipient="Главному врачу\nГБУЗ «Раменская больница»",
        postal_address=(
            "140100, Московская область, г. Раменское, ул. Махова, д. 14"
        ),
        phones=("+7 (496) 463-30-21",),
        email="kanc@ramcrb.ru",
        verification_status="user_verified",
    )

    output_path = GenerationService().generate(
        person=PersonFactory.create(),
        organization=organization,
        template_path=template_path,
        output_dir=tmp_path / "output",
    )

    generated = Document(output_path)
    assert generated.paragraphs[0].runs[0].font.highlight_color is None


def test_generation_simplifies_recipient_header():
    organization = Organization(
        id=None,
        organization_type=OrganizationType.HOSPITAL,
        recipient=(
            "Главному врачу\n"
            "ГБУЗ «Раменская больница» | Клиники в Раменском\n"
            "полковнику полиции\n"
            "В.Н. Савкину"
        ),
        postal_address=(
            "140100, Московская область, г. Раменское. "
            "ИНН: 1234567890; Часы работы: 08:00–17:00"
        ),
        verification_status="auto_found",
    )

    header = GenerationService._simple_recipient_header(organization)

    assert "ГБУЗ «Раменская больница»" in header
    assert "140100, Московская область, г. Раменское" in header
    assert "Клиники в Раменском" not in header
    assert "В.Н. Савкину" not in header
    assert "полковнику полиции" not in header
    assert "ИНН" not in header
    assert "Часы работы" not in header

    noisy_title = Organization(
        id=None,
        organization_type=OrganizationType.ELECTION_COMMISSION,
        recipient=(
            "Председателю\n"
            "ТИК Раменское - сайт, телефон, режим работы ТИК города Раменское"
        ),
        postal_address="7702129350) Московская область",
        verification_status="auto_found",
    )
    cleaned = GenerationService._simple_recipient_header(noisy_title)
    assert "сайт, телефон" not in cleaned
    assert "7702129350" not in cleaned
    assert "ТИК" not in cleaned
    assert "Территориальной избирательной комиссии" in cleaned

    directory_result = Organization(
        id=None,
        organization_type=OrganizationType.PSYCHIATRY,
        recipient="Областная психиатрическая больница... : GosRegion",
        postal_address="г. Раменское, ул. Больничная, д. 6 Часы работы",
        verification_status="auto_found",
    )
    cleaned = GenerationService._simple_recipient_header(directory_result)
    assert "GosRegion" not in cleaned
    assert "Часы работы" not in cleaned

    incomplete = Organization(
        id=None,
        organization_type=OrganizationType.HOSPITAL,
        recipient="Главному врачу\nРаменской больницы",
        postal_address="ПОЧТОВЫЙ АДРЕС ТРЕБУЕТ ПРОВЕРКИ",
        phones=("+7 495 111-11-11", "+7 495 222-22-22", "+7 495 333-33-33"),
        email="ТРЕБУЕТ ПРОВЕРКИ",
        verification_status="needs_review",
    )
    cleaned = GenerationService._simple_recipient_header(incomplete)
    assert "ТРЕБУЕТ ПРОВЕРКИ" not in cleaned
    assert "+7 495 333-33-33" not in cleaned


def test_generation_abbreviates_long_medical_titles():
    organization = Organization(
        id=None,
        organization_type=OrganizationType.HOSPITAL,
        recipient=(
            "Главному врачу\n"
            "Государственное бюджетное учреждение здравоохранения "
            "Московской области «Клиническая областная больница»"
        ),
        postal_address="141300, Московская область, г. Сергиев Посад",
        phones=("+7 (496) 111-11-11",),
        email="hospital@example.ru",
    )

    header = GenerationService._simple_recipient_header(organization)

    assert "ГБУЗ Московской области «КОБ»" in header
    assert "Государственное бюджетное учреждение" not in header


def test_dynamic_recipient_header_uses_tnr_and_keeps_template_size():
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("{{RECIPIENT_HEADER}}")
    run.font.name = "Arial"
    run.font.size = Pt(13)

    GenerationService._replace_in_paragraph(
        paragraph,
        {"{{RECIPIENT_HEADER}}": "Главному врачу\nГБУЗ «ЦРБ»"},
    )

    assert paragraph.runs[0].font.name == "Times New Roman"
    assert paragraph.runs[0].font.size == Pt(13)


def test_generation_uses_one_blank_line_before_contiguous_contacts():
    organization = Organization(
        id=None,
        organization_type=OrganizationType.ELECTION_COMMISSION,
        recipient=(
            "Председателю\nТерриториальной избирательной комиссии\n"
            "Раменского городского округа"
        ),
        postal_address=(
            "140100, Московская область, г. Раменское, "
            "Комсомольская площадь, д. 2, каб. 116"
        ),
        phones=("8 (496) 463-67-03",),
        email="tikramenskoe@mail.ru",
    )

    header = GenerationService._simple_recipient_header(organization)

    assert header == (
        "Председателю\n"
        "Территориальной избирательной комиссии\n"
        "Раменского городского округа\n\n"
        "140100, Московская область, г. Раменское, "
        "Комсомольская площадь, д. 2, каб. 116\n"
        "тел.: 8 (496) 463-67-03\n"
        "эл. почта: tikramenskoe@mail.ru"
    )


def test_generation_keeps_found_email_for_protocol_exceptions():
    organization = Organization(
        id=None,
        organization_type=OrganizationType.MILITARY_COMMANDANT,
        recipient="Военному коменданту\nВоенной комендатуры\nПодольского гарнизона",
        postal_address="142100, Московская область, г. Подольск, ул. Парковая, д. 56",
        phones=("+7 (496) 755-63-30",),
        email="commandant@example.ru",
    )

    header = GenerationService._simple_recipient_header(organization)

    assert "эл. почта: commandant@example.ru" in header


def test_generation_removes_repeated_military_organization_names():
    commissariat = Organization(
        id=None,
        organization_type=OrganizationType.MILITARY_COMMISSARIAT,
        recipient=(
            "Военному комиссару\n"
            "Военного комиссариата\n"
            "Объединённый военный комиссариат\n"
            "Черёмушкинского района Юго-Западного административного округа "
            "города Москвы"
        ),
        postal_address="119333, г. Москва, ул. Вавилова, д. 44, к. 1",
    )
    commandant = Organization(
        id=None,
        organization_type=OrganizationType.MILITARY_COMMANDANT,
        recipient=(
            "Военному коменданту\n"
            "Военной комендатуры\n"
            "Военная комендатура Челябинского гарнизона"
        ),
        postal_address="454091, г. Челябинск, ул. Кирова, д. 92",
    )

    commissariat_header = GenerationService._simple_recipient_header(commissariat)
    commandant_header = GenerationService._simple_recipient_header(commandant)

    assert commissariat_header.startswith(
        "Военному комиссару\n"
        "Военного комиссариата\n"
        "Черёмушкинского района Юго-Западного административного округа "
        "города Москвы\n\n"
    )
    assert "Объединённый военный комиссариат" not in commissariat_header
    assert commandant_header.startswith(
        "Военному коменданту\n"
        "Военной комендатуры\n"
        "Челябинского гарнизона\n\n"
    )
    assert commandant_header.count("комендатур") == 1


def test_generation_formats_contract_service_point_as_four_fixed_lines():
    organization = Organization(
        id=None,
        organization_type=OrganizationType.CONTRACT_SERVICE_POINT,
        recipient=(
            "Начальнику\nПункт отбора на военную службу по контракту»\n"
            "Московская область"
        ),
        postal_address=(
            "143900, Московская область, г. Балашиха, "
            "Восточное шоссе, владение 2"
        ),
        phones=("+7 (495) 990-77-77",),
        email="unused@example.ru",
    )

    header = GenerationService._simple_recipient_header(organization)

    assert header.startswith(
        "Начальнику\nПункта отбора\nна военную службу по контракту\n"
        "Московской области\n\n"
    )
    assert "эл. почта: unused@example.ru" in header


def test_generation_removes_administration_head_and_word_administration():
    organization = Organization(
        id=None,
        organization_type=OrganizationType.ADMINISTRATION,
        recipient=(
            "Главе\nАдминистрации Раменского муниципального округа\n"
            "Малышеву Эдуарду Владимировичу"
        ),
        postal_address=(
            "140100, Московская область, г. Раменское, "
            "Комсомольская площадь, д. 2"
        ),
        phones=("+7 (496) 473-91-01",),
        email="ram_glava@mosreg.ru",
    )

    header = GenerationService._simple_recipient_header(organization)

    assert header.startswith(
        "Главе\nРаменского муниципального округа\n\n"
    )
    assert "Малышеву Эдуарду Владимировичу" not in header


def test_generation_enforces_canonical_dynamic_header_labels():
    commissariat = Organization(
        id=None,
        organization_type=OrganizationType.MILITARY_COMMISSARIAT,
        recipient="Военному комиссару\nРаменского района\nМосковской области",
        postal_address="140100, г. Раменское, ул. Гурьева, д. 21",
    )
    commandant = Organization(
        id=None,
        organization_type=OrganizationType.MILITARY_COMMANDANT,
        recipient="Военному коменданту\nНОГИНСКОГО ГАРНИЗОНА",
        postal_address="Московская область, г. Ногинск",
    )
    tfoms = Organization(
        id=None,
        organization_type=OrganizationType.TFOMS,
        recipient="Директору\nТФОМС «МО» Московская область",
        postal_address="Московская область",
    )

    commissariat_header = GenerationService._simple_recipient_header(commissariat)
    commandant_header = GenerationService._simple_recipient_header(commandant)
    tfoms_header = GenerationService._simple_recipient_header(tfoms)

    assert "Военному комиссару\nВоенного комиссариата\nРаменского района" in commissariat_header
    assert "Военному коменданту\nВоенной комендатуры\nНогинского гарнизона" in commandant_header
    assert "Территориального фонда\nобязательного медицинского страхования" in tfoms_header
    assert "Московской области" in tfoms_header
    assert "ТФОМС" not in tfoms_header
    assert "\n\nМосковская область" not in tfoms_header


def test_generation_creates_complete_fifteen_document_bundle(tmp_path):
    template_dir = tmp_path / "templates"
    template_dir.mkdir()
    for request in REQUEST_CATALOG:
        template = Document()
        template.add_paragraph(
            f"{request.number}: {{{{RECIPIENT_HEADER}}}} "
            "{{LAST_NAME_GENITIVE}} {{INITIALS}}"
        )
        template.save(template_dir / request.filename)

    organizations = [
        Organization(
            id=index,
            organization_type=organization_type,
            recipient=f"Адресат {index}",
            postal_address=f"Адрес {index}",
            verification_status="test",
        )
        for index, organization_type in enumerate(
            (
                request.organization_type
                for request in REQUEST_CATALOG
                if request.organization_type is not None
            ),
            start=1,
        )
    ]

    generated = GenerationService().generate_bundle(
        person=PersonFactory.create(),
        organizations=organizations,
        template_dir=template_dir,
        output_dir=tmp_path / "output",
    )

    assert len(generated) == 15
    assert [path.name[:2] for path in generated] == [
        f"{number:02d}" for number in range(1, 16)
    ]
    assert all(path.is_file() for path in generated)


def test_editable_bundle_keeps_request_bookmarks(tmp_path):
    sources = []
    for request in REQUEST_CATALOG:
        source = tmp_path / request.filename
        document = Document()
        document.add_paragraph(f"Запрос №{request.number}")
        document.save(source)
        sources.append(source)

    output = tmp_path / "combined.docx"
    WordBundleService().create_editable_bundle(sources, output)

    combined = Document(output)
    bookmark_names = {
        element.get(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name"
        )
        for element in combined.element.body.iter()
        if element.tag.endswith("bookmarkStart")
    }
    assert len(combined.sections) == 15
    assert bookmark_names == {
        f"NEXUSDOCS_REQUEST_{number:02d}" for number in range(1, 16)
    }


def test_editable_bundle_normalizes_footer_number_on_every_section(tmp_path):
    sources = []
    section_counts = []
    for request in REQUEST_CATALOG:
        source = tmp_path / request.filename
        document = Document()
        document.add_paragraph(f"Запрос №{request.number}")
        footer = document.sections[0].footer
        footer.paragraphs[0].text = f"Исх. № 5002/99-26"
        section_count = 1
        if request.number in {9, 11}:
            continuation = document.add_section(WD_SECTION.NEW_PAGE)
            continuation.footer.is_linked_to_previous = False
            continuation.footer.paragraphs[0].text = ""
            section_count = 2
        document.save(source)
        sources.append(source)
        section_counts.append(section_count)

    output = tmp_path / "combined.docx"
    WordBundleService().create_editable_bundle(sources, output)
    combined = Document(output)

    actual_numbers = []
    for section in combined.sections:
        footer_text = "\n".join(
            [paragraph.text for paragraph in section.footer.paragraphs]
            + [
                paragraph.text
                for table in section.footer.tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            ]
        )
        match = __import__("re").search(r"/(\d{1,2})-26", footer_text)
        actual_numbers.append(int(match.group(1)) if match else None)

    expected_numbers = [
        (
            None
            if index and request.number in {9, 11}
            else request.number
        )
        for request, count in zip(REQUEST_CATALOG, section_counts, strict=True)
        for index in range(count)
    ]
    assert actual_numbers == expected_numbers
