from pathlib import Path

from docx import Document

from domain.entities.organization import Organization
from domain.enums.organization_type import OrganizationType
from services.generation_service import GenerationService
from services.word_typography_service import protect_word_line_breaks
from tests.factories.person_factory import PersonFactory


def test_protects_general_articles_and_military_units_across_runs(tmp_path: Path):
    path = tmp_path / "typography.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Ответ направить по адресу. В отношении в/часть ")
    bold = paragraph.add_run("01007, ч. 2.1 ст.")
    bold.bold = True
    paragraph.add_run(" 337 УК РФ применяется. ")
    paragraph.add_run(
        "ч. 5 ст. 337 УК РФ и ст. 338 УК РФ. "
        "СК России РФ действует. УК РФ действует. "
        "пос. Власиха, 10 суток. ПО дороге"
    )
    document.save(path)

    protect_word_line_breaks(path)
    result = Document(path)
    text = result.paragraphs[0].text
    assert "по адресу" in text
    assert "В отношении" in text
    assert "в/часть\u00a001007" in text
    assert "ч.\u00a02.1\u00a0ст.\u00a0337\u00a0УК\u00a0РФ применяется" in text
    assert "ч.\u00a05\u00a0ст.\u00a0337\u00a0УК\u00a0РФ" in text
    assert "ст.\u00a0338\u00a0УК\u00a0РФ" in text
    assert "СК России РФ действует" in text
    assert "УК РФ действует" in text
    assert "пос.\u00a0Власиха" in text
    assert "10 суток" in text
    assert "ПО дороге" in text
    assert result.paragraphs[0].runs[1].bold is True

    first_pass = text
    protect_word_line_breaks(path)
    assert Document(path).paragraphs[0].text == first_pass


def test_does_not_join_across_explicit_breaks_or_paragraphs(tmp_path: Path):
    path = tmp_path / "breaks.docx"
    document = Document()
    paragraph = document.add_paragraph("на")
    paragraph.add_run().add_break()
    paragraph.add_run("службу")
    document.add_paragraph("в")
    document.add_paragraph("часть")
    document.save(path)

    protect_word_line_breaks(path)
    result = Document(path)
    assert [paragraph.text for paragraph in result.paragraphs] == [
        "на\nслужбу", "в", "часть",
    ]


def test_generation_protects_dynamic_article_after_placeholder_replacement(
    tmp_path: Path,
):
    template = tmp_path / "request.docx"
    document = Document()
    document.add_paragraph("По делу военнослужащего в/часть {{MILITARY_UNIT}}")
    document.add_paragraph("Ответ по статье {{ARTICLE}} УК РФ")
    document.save(template)

    organization = Organization(
        id=1,
        organization_type=OrganizationType.MILITARY_COMMISSARIAT,
        recipient="Военному комиссару",
        postal_address="140100, Московская область",
        phones=("+7 (496) 000-00-00",),
        email="example@example.ru",
        verification_status="test",
    )
    output = GenerationService().generate(
        PersonFactory.create(), organization, template, tmp_path / "out",
    )
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "в/часть\u00a0" in text
    assert "УК РФ" in text
