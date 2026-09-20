from __future__ import annotations

from dataclasses import replace
from datetime import date
from pathlib import Path

from docx import Document

from services.extension_generation_service import ExtensionGenerationService
from tests.factories.person_factory import PersonFactory


TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1]
    / "templates"
    / "extensions"
    / "extension_to_10_days.docx"
)


def _all_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def test_generates_extension_with_calculated_dates_and_no_times(tmp_path):
    person = PersonFactory.create()
    person = replace(
        person,
        military=replace(
            person.military,
            military_unit="в/ч 95375",
            military_deployment="г. Ясиноватая ДНР",
            service_basis="по мобилизации",
            rank="старшина",
            position="заместитель командира взвода",
        ),
        soch_case=replace(
            person.soch_case,
            soch_date=date(2026, 8, 10),
            registration_date=date(2026, 7, 25),
            circumstances="Неявка в срок",
            article="ч. 2.1 ст. 337 УК РФ",
        ),
        document_info=replace(
            person.document_info,
            outgoing_date=date(2026, 8, 17),
        ),
    )

    output = ExtensionGenerationService().generate(
        person,
        TEMPLATE_PATH,
        tmp_path,
    )
    text = _all_text(output)

    assert output.name == "ИвановИИ_17.08.2026_продление_до_10_суток.docx"
    assert "17.08.2026" in text
    assert "19.08.2026" in text
    assert "26.08.2026" in text
    assert "25.07.2026" not in text
    assert "по мобилизации в в/части 95375" in text
    assert "Иванова Ивана Ивановича" in text
    assert "старшины Иванова И.И." in text
    assert "ч. 2.1 ст. 337 УК РФ" in text
    assert "совершил неявку в срок без уважительных причин" in text
    assert "Около 00 часов" not in text
    assert "Около 15 часов" not in text
    assert "{{" not in text
    assert "Шмел" not in text


def test_generates_second_supported_circumstance(tmp_path):
    person = PersonFactory.create()
    person = replace(
        person,
        soch_case=replace(
            person.soch_case,
            circumstances="Самовольное оставление части",
        ),
    )

    output = ExtensionGenerationService().generate(
        person,
        TEMPLATE_PATH,
        tmp_path,
    )

    assert "совершил самовольное оставление в/части" in _all_text(output)


def test_reports_all_new_fields_missing_for_legacy_person():
    person = PersonFactory.create()
    person = replace(
        person,
        military=replace(
            person.military,
            military_deployment="",
            service_basis="",
        ),
        soch_case=replace(
            person.soch_case,
            registration_date=None,
            circumstances="",
        ),
    )

    assert ExtensionGenerationService.missing_fields(person) == (
        "Дислокация в/части",
        "Условия призыва",
        "Обстоятельства СОЧ",
    )


def test_extension_dates_use_outgoing_date_for_legacy_person():
    person = PersonFactory.create()
    person = replace(
        person,
        soch_case=replace(person.soch_case, registration_date=None),
        document_info=replace(
            person.document_info,
            outgoing_date=date(2026, 12, 30),
        ),
    )

    replacements = ExtensionGenerationService._build_replacements(person)

    assert ExtensionGenerationService.missing_fields(person) == ()
    assert replacements["{{REGISTRATION_DATE}}"] == "30.12.2026"
    assert replacements["{{THREE_DAY_DATE}}"] == "01.01.2027"
    assert replacements["{{TEN_DAY_DATE}}"] == "08.01.2027"
    assert "30.12.2026" in ExtensionGenerationService.filename(person)
