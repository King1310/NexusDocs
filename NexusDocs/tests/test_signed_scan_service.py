from __future__ import annotations

from datetime import date
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter
import pytest

from services.signed_scan_service import (
    DISPATCH_RULES,
    DispatchPackageStore,
    RecognizedScanRequest,
    ScanRequestRecognizer,
    SignedScanService,
    WordBundleHeaderEmails,
    WordBundlePageMap,
)
from tests.factories.person_factory import PersonFactory


class StaticRequestRecognizer:
    @staticmethod
    def recognize(scan_path: Path, outgoing_number: str):
        return (
            RecognizedScanRequest(13, 0, 2, ""),
            RecognizedScanRequest(5, 2, 1, "scan5@example.ru"),
            RecognizedScanRequest(16, 3, 2, "second@example.ru"),
        )


class StaticHeaderEmails:
    @staticmethod
    def emails_by_request(bundle_path: Path) -> dict[int, str]:
        return {
            rule.container_number: (
                ""
                if rule.container_number == 13
                else f"request{rule.container_number}@example.ru"
            )
            for rule in DISPATCH_RULES
        }


class AllContainersRecognizer:
    @staticmethod
    def recognize(scan_path: Path, outgoing_number: str):
        requests = []
        first_page = 0
        for container_number in range(1, 16):
            page_count = 2 if container_number == 9 else 1
            requests.append(
                RecognizedScanRequest(
                    container_number,
                    first_page,
                    page_count,
                    "",
                )
            )
            first_page += page_count
        return tuple(requests)


class HeaderEmailsMustNotBeRead:
    @staticmethod
    def emails_by_request(bundle_path: Path) -> dict[int, str]:
        raise AssertionError("Word e-mail fallback must stay optional")


class RepeatedAndForeignRecognizer:
    @staticmethod
    def recognize(scan_path: Path, outgoing_number: str):
        return (
            RecognizedScanRequest(
                6,
                0,
                1,
                "first-psych@example.ru",
                "5771.6",
                date(2026, 9, 10),
                0,
            ),
            RecognizedScanRequest(
                6,
                1,
                1,
                "second-psych@example.ru",
                "5771.6",
                date(2026, 9, 10),
                1,
            ),
            RecognizedScanRequest(
                0,
                2,
                2,
                "other@example.ru",
                "5891",
                date(2026, 9, 11),
                0,
            ),
        )


def _pdf(path: Path, page_count: int) -> Path:
    writer = PdfWriter()
    for index in range(page_count):
        writer.add_blank_page(width=595 + index, height=842)
    with path.open("wb") as stream:
        writer.write(stream)
    return path


def _bookmark(paragraph, request_number: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(1000 + request_number))
    start.set(qn("w:name"), f"NEXUSDOCS_REQUEST_{request_number:02d}")
    paragraph._p.insert(0, start)


def test_dispatch_rules_include_only_the_ten_scanned_requests():
    assert [rule.container_number for rule in DISPATCH_RULES] == [
        3,
        4,
        5,
        6,
        7,
        8,
        12,
        13,
        14,
        15,
    ]
    assert [rule.abbreviation for rule in DISPATCH_RULES] == [
        "комиссариат",
        "комендатура",
        "нарк",
        "псих",
        "црб",
        "тфомс",
        "администрация",
        "повск",
        "фку",
        "тик",
    ]


def test_signed_scan_uses_outgoing_suffix_instead_of_pdf_position(tmp_path):
    person = PersonFactory.create()
    person.document_info = type(person.document_info)(
        outgoing_number="4755",
        outgoing_date=date(2026, 9, 12),
    )
    scan = _pdf(tmp_path / "scan.pdf", 5)
    bundle = tmp_path / "bundle.docx"
    bundle.touch()
    output = tmp_path / "ready"

    result = SignedScanService(
        request_recognizer=StaticRequestRecognizer(),
        header_emails=StaticHeaderEmails(),
    ).prepare(
        person=person,
        bundle_path=bundle,
        scan_path=scan,
        output_directory=output,
    )

    assert len(result.attachments) == 3
    assert result.attachments[0].path.name == (
        "исх. № 4755.13 от 12.09.2026 повск.pdf"
    )
    assert result.attachments[1].path.name == (
        "исх. № 4755.5 от 12.09.2026 нарк.pdf"
    )
    assert result.attachments[1].subject == (
        "Запрос в отношении Иванова И.И. нарк"
    )
    assert result.attachments[2].path.name == (
        "исх. № 4755.16 от 12.09.2026.pdf"
    )
    assert result.attachments[2].subject == (
        "Запрос в отношении Иванова И.И."
    )
    assert len(PdfReader(result.attachments[0].path).pages) == 2
    assert len(PdfReader(result.attachments[-1].path).pages) == 2
    assert result.attachments[0].recipient_email == ""
    assert result.attachments[1].recipient_email == "scan5@example.ru"
    assert result.attachments[2].recipient_email == "second@example.ru"
    assert result.missing_email_attachments == (result.attachments[0],)
    assert result.manifest_path.name == "пакет_отправки.nexusdocs.json"
    assert result.manifest_path.is_file()
    launcher = result.output_directory / DispatchPackageStore.LAUNCHER_FILENAME
    powershell = (
        result.output_directory / DispatchPackageStore.POWERSHELL_FILENAME
    )
    assert launcher.is_file()
    assert powershell.is_file()
    launcher_text = launcher.read_text(encoding="ascii")
    powershell_text = powershell.read_text(encoding="utf-8-sig")
    assert "SysWOW64" in launcher_text
    assert "Outlook.Application" in powershell_text
    assert "$mail.Attachments.Add(" in powershell_text
    assert "$displayName" in powershell_text
    assert "$entry.Item.outgoing_number" in powershell_text
    assert "$entry.Item.duplicate_index" in powershell_text
    assert "$outlook.Explorers.Add($drafts, 0)" in powershell_text
    assert "$mail.Save()" in powershell_text
    assert "$mail.Send()" not in powershell_text

    copied_directory = tmp_path / "copied_to_desktop"
    shutil.copytree(result.output_directory, copied_directory)
    copied = DispatchPackageStore().load(
        copied_directory / result.manifest_path.name
    )

    assert copied.output_directory == copied_directory.resolve()
    assert (
        copied_directory / DispatchPackageStore.LAUNCHER_FILENAME
    ).is_file()
    assert [item.path.name for item in copied.attachments] == [
        item.path.name for item in result.attachments
    ]
    assert [item.subject for item in copied.attachments] == [
        item.subject for item in result.attachments
    ]
    assert copied.attachments[0].recipient_email == ""


def test_request_recognizer_groups_repeated_markers_and_arbitrary_order():
    requests = ScanRequestRecognizer.requests_from_page_texts(
        (
            "08.09.2026 № 5800/13 эл. почта: first@example.ru",
            "Исх. № Исх-5800/13-26",
            "08.09.2026 № 5800.5 эл. почта: narc@example.ru",
            "08.09.2026 № 5800/16 эл. почта: second@example.ru",
            "Продолжение последнего запроса",
        ),
        "5800",
    )

    assert [request.container_number for request in requests] == [13, 5, 16]
    assert [request.page_count for request in requests] == [2, 1, 2]
    assert requests[0].recipient_email == "first@example.ru"
    assert requests[2].recipient_email == "second@example.ru"


def test_request_recognizer_numbers_repeated_container_in_any_position():
    requests = ScanRequestRecognizer.requests_from_page_texts(
        (
            "08.09.2026 № 5800/13 В производстве военного "
            "следственного отдела",
            "08.09.2026 № 5800/5 В производстве военного "
            "следственного отдела",
            "08.09.2026 № 5800/13 В производстве военного "
            "следственного отдела",
        ),
        "5800",
    )

    assert [request.container_number for request in requests] == [13, 5, 13]
    assert [request.duplicate_index for request in requests] == [0, 0, 1]
    assert [request.page_count for request in requests] == [1, 1, 1]


def test_request_recognizer_splits_consecutive_duplicate_first_pages():
    requests = ScanRequestRecognizer.requests_from_page_texts(
        (
            "08.09.2026 № 5771/6 В производстве военного "
            "следственного отдела эл. почта: first@example.ru",
            "08.09.2026 № 5771/6 В производстве военного "
            "следственного отдела эл. почта: second@example.ru",
            "08.09.2026 № 5771/6 В производстве военного "
            "следственного отдела",
        ),
        "5771",
    )

    assert [request.duplicate_index for request in requests] == [0, 1, 2]
    assert [request.page_count for request in requests] == [1, 1, 1]
    assert [request.recipient_email for request in requests] == [
        "first@example.ru",
        "second@example.ru",
        "",
    ]


def test_request_recognizer_accepts_joined_ocr_number():
    requests = ScanRequestRecognizer.requests_from_page_texts(
        (
            "07.09.2026 № 57776 В производстве военного "
            "следственного отдела",
            "07.09.2026 № 57778 В производстве военного "
            "следственного отдела",
        ),
        "5777",
    )

    assert [request.outgoing_number for request in requests] == [
        "5777.6",
        "5777.8",
    ]
    assert [request.container_number for request in requests] == [6, 8]


def test_request_recognizer_keeps_foreign_outgoing_and_its_continuation():
    requests = ScanRequestRecognizer.requests_from_page_texts(
        (
            "08.09.2026 № 5771/5 В производстве военного "
            "следственного отдела",
            "11.09.2026 № 5891 Другой запрос "
            "эл. почта: other@example.ru",
            "Продолжение другого запроса без исходящего номера",
            "08.09.2026 № 5771/7 В производстве военного "
            "следственного отдела",
        ),
        "5771",
    )

    assert [request.outgoing_number for request in requests] == [
        "5771.5",
        "5891",
        "5771.7",
    ]
    assert [request.container_number for request in requests] == [5, 0, 7]
    assert [request.page_count for request in requests] == [1, 2, 1]
    assert requests[1].recipient_email == "other@example.ru"
    assert requests[1].outgoing_date == date(2026, 9, 11)


def test_repeated_and_foreign_outgoing_filenames_and_portable_package(tmp_path):
    person = PersonFactory.create()
    person.document_info = type(person.document_info)(
        outgoing_number="5771",
        outgoing_date=date(2026, 9, 10),
    )
    scan = _pdf(tmp_path / "mixed.pdf", 4)

    result = SignedScanService(
        request_recognizer=RepeatedAndForeignRecognizer(),
        header_emails=HeaderEmailsMustNotBeRead(),
    ).prepare(
        person=person,
        bundle_path=None,
        scan_path=scan,
        output_directory=tmp_path / "mixed_split",
    )

    assert [attachment.path.name for attachment in result.attachments] == [
        "исх. № 5771.6 от 10.09.2026 псих.pdf",
        "исх. № 5771.6.1 от 10.09.2026 псих.pdf",
        "исх. № 5891 от 11.09.2026.pdf",
    ]
    assert [attachment.subject for attachment in result.attachments] == [
        "Запрос в отношении Иванова И.И. псих",
        "Запрос в отношении Иванова И.И. псих",
        "Запрос в отношении Иванова И.И.",
    ]
    copied = DispatchPackageStore().load(result.manifest_path)
    assert [attachment.container_number for attachment in copied.attachments] == [
        6,
        6,
        0,
    ]
    assert [
        attachment.display_outgoing_number
        for attachment in copied.attachments
    ] == ["5771.6", "5771.6.1", "5891"]


def test_any_container_is_split_and_word_bundle_is_optional(tmp_path):
    person = PersonFactory.create()
    person.document_info = type(person.document_info)(
        outgoing_number="4755",
        outgoing_date=date(2026, 9, 12),
    )
    scan = _pdf(tmp_path / "all_requests.pdf", 16)

    result = SignedScanService(
        request_recognizer=AllContainersRecognizer(),
        header_emails=HeaderEmailsMustNotBeRead(),
    ).prepare(
        person=person,
        bundle_path=None,
        scan_path=scan,
        output_directory=tmp_path / "all_split",
    )

    assert [item.container_number for item in result.attachments] == list(
        range(1, 16)
    )
    ninth = result.attachments[8]
    assert ninth.container_number == 9
    assert ninth.abbreviation == ""
    assert ninth.path.name == "исх. № 4755.9 от 12.09.2026.pdf"
    assert len(PdfReader(ninth.path).pages) == 2
    assert result.attachments[4].path.name.endswith(" нарк.pdf")


def test_dispatch_package_detects_changed_pdf(tmp_path):
    person = PersonFactory.create()
    scan = _pdf(tmp_path / "scan.pdf", 5)
    bundle = tmp_path / "bundle.docx"
    bundle.touch()
    result = SignedScanService(
        request_recognizer=StaticRequestRecognizer(),
        header_emails=StaticHeaderEmails(),
    ).prepare(
        person=person,
        bundle_path=bundle,
        scan_path=scan,
        output_directory=tmp_path / "ready",
    )
    result.attachments[0].path.write_bytes(b"changed")

    with pytest.raises(ValueError, match="изменён или повреждён"):
        DispatchPackageStore().load(result.manifest_path)


def test_header_email_reader_ignores_the_response_address(tmp_path):
    document = Document()
    document.add_paragraph("эл. почта: commissariat@example.ru")
    request_three = document.add_paragraph()
    _bookmark(request_three, 3)

    document.add_section(WD_SECTION.NEW_PAGE)
    request_four = document.add_paragraph()
    _bookmark(request_four, 4)
    document.add_paragraph(
        "Ответ направить электронной почтой по адресу: vso@example.ru"
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Электронная почта: fku@example.ru")
    request_fourteen = document.add_paragraph()
    _bookmark(request_fourteen, 14)
    path = tmp_path / "bundle.docx"
    document.save(path)

    emails = WordBundleHeaderEmails().emails_by_request(path)

    assert emails[3] == "commissariat@example.ru"
    assert emails[4] == ""
    assert emails[14] == "fku@example.ru"


def test_page_counts_are_calculated_from_bookmark_starts():
    starts = {number: number for number in range(1, 16)}
    starts[4] = 5
    for number in range(5, 16):
        starts[number] = number + 1

    counts = WordBundlePageMap._page_counts_from_starts(starts, 17)

    assert counts[3] == 2
    assert counts[4] == 1
    assert counts[15] == 2
